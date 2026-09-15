from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ENTITY_PURPOSES = {
    "WarehouseAttendancePolicy": "Bảng WarehouseAttendancePolicy lưu cấu hình chấm công áp dụng cho một cơ sở, gồm phương thức xác minh, phạm vi IP/GPS, các chế độ làm việc được phép và thời gian hiệu lực của chính sách.",
    "AttendanceWorkArrangement": "Bảng AttendanceWorkArrangement lưu các thỏa thuận làm việc ngoài cơ sở của nhân viên, như làm việc tại nhà hoặc đi công tác, để hệ thống áp dụng đúng quy tắc vị trí khi chấm công trong khoảng thời gian đã được phê duyệt.",
    "WarehouseAttendanceExemption": "Bảng WarehouseAttendanceExemption lưu ngoại lệ chấm công của từng người dùng tại một cơ sở, cho biết người đó có bắt buộc chấm công hay không trong thời gian hiệu lực.",
    "AttendanceLog": "Bảng AttendanceLog lưu kết quả chấm công hằng ngày của nhân viên, gồm thời điểm chấm công, cơ sở, phương thức xác minh, dữ liệu vị trí, trạng thái chấp nhận hoặc từ chối và lý do từ chối.",
    "AttendanceLateReport": "Bảng AttendanceLateReport lưu báo cáo đi muộn do nhân viên gửi, thời gian dự kiến đến, lý do, bản ghi chấm công liên quan và kết quả người có thẩm quyền xem xét.",
    "EmployeeContract": "Bảng EmployeeContract lưu thông tin hợp đồng lao động của nhân viên, gồm số hợp đồng, loại hợp đồng, thời hạn, trạng thái, quan hệ gia hạn và thông tin chấm dứt hoặc hủy hợp đồng.",
    "EmployeeContractAutomationRun": "Bảng EmployeeContractAutomationRun ghi nhận từng lần chạy tác vụ tự động liên quan đến hợp đồng, phục vụ theo dõi trạng thái xử lý, số lần thử, kết quả và lỗi phát sinh.",
    "EmployeeContractDocument": "Bảng EmployeeContractDocument lưu thông tin tệp hợp đồng đã hoàn tất tải lên, gồm vị trí lưu trữ, mã kiểm tra nội dung, phiên bản và người thực hiện.",
    "EmployeeContractDocumentUploadIntent": "Bảng EmployeeContractDocumentUploadIntent lưu yêu cầu tải tệp hợp đồng trước khi tệp được xác nhận hoàn tất, giúp kiểm soát loại tệp, dung lượng, đường dẫn đích, thời hạn và kết quả xử lý.",
    "EmployeeContractImportRow": "Bảng EmployeeContractImportRow lưu kết quả xử lý từng dòng trong một lô nhập hợp đồng, gồm dữ liệu đã chuẩn hóa, lỗi kiểm tra, tài liệu đính kèm và bản ghi được tạo sau khi ghi nhận thành công.",
    "EmployeeProfile": "Bảng EmployeeProfile lưu hồ sơ nhân sự dùng chung cho các nghiệp vụ nhân viên, gồm mã nhân viên, thông tin liên hệ, chức danh, phòng ban, cơ sở làm việc và trạng thái làm việc.",
    "EmployeeEmploymentTransition": "Bảng EmployeeEmploymentTransition lưu yêu cầu và lịch sử chuyển trạng thái làm việc của nhân viên, ví dụ từ thử việc sang chính thức hoặc nghỉ việc, cùng ngày hiệu lực và người xử lý.",
    "ExternalScanQueue": "Bảng ExternalScanQueue lưu các lượt quét hàng từ thiết bị hoặc hệ thống bên ngoài trước khi được phê duyệt và chuyển thành chứng từ xuất, đồng thời hỗ trợ theo dõi ca làm việc, bàn giao, giữ tồn khả dụng và đồng bộ ngoại tuyến.",
    "ExternalQueueScannableProductConfig": "Bảng ExternalQueueScannableProductConfig cấu hình danh sách sản phẩm được phép quét từ nguồn bên ngoài tại từng cơ sở hoặc khu vực lưu trữ.",
    "OfficeScopeConfig": "Bảng OfficeScopeConfig lưu chính sách phạm vi truy cập của một văn phòng/cơ sở, xác định chế độ tính phạm vi, phiên bản chính sách và thời gian hiệu lực.",
    "OfficeScopeEdge": "Bảng OfficeScopeEdge lưu quan hệ cho phép một văn phòng/cơ sở truy cập đến một cơ sở đích, dùng làm đầu vào tính quyền truy cập theo phạm vi.",
    "OfficeScopeMaterialization": "Bảng OfficeScopeMaterialization theo dõi một lần tính và tạo trước dữ liệu phạm vi truy cập của văn phòng/cơ sở, gồm tiến độ, số bản ghi thành công hoặc thất bại và lỗi gần nhất.",
    "UserFacilityAccessGrant": "Bảng UserFacilityAccessGrant lưu quyền truy cập đã được tính cho một người dùng tại một cơ sở cụ thể, bao gồm loại cơ sở, tập quyền và nguồn hình thành quyền.",
    "UserAccessVersion": "Bảng UserAccessVersion lưu một phiên bản đầy đủ của quyền truy cập đã tính cho người dùng, giúp kích hoạt phiên bản mới, lưu vết phiên bản cũ và kiểm soát thay đổi chính sách.",
    "UserAccessMetadata": "Bảng UserAccessMetadata lưu trạng thái tổng hợp hiện hành của quyền truy cập người dùng, trỏ đến phiên bản đang hoạt động và các thông tin dùng để phát hiện khi cần tính lại quyền.",
    "FileTemplate": "Bảng FileTemplate lưu danh mục tệp biểu mẫu dùng trong hệ thống, gồm tên, nhóm, định dạng, đường dẫn tải xuống và lịch sử phiên bản.",
    "Inventory": "Bảng Inventory lưu số lượng tồn hiện tại của một sản phẩm tại một cơ sở hoặc khu vực lưu trữ, tách rõ tổng tồn, tồn khả dụng, tồn đang giữ, đang vận chuyển và đang cách ly.",
    "InventoryStockPolicy": "Bảng InventoryStockPolicy lưu ngưỡng tồn kho theo phạm vi áp dụng, dùng để cảnh báo tồn tối thiểu, giới hạn tồn tối đa và đề xuất thời điểm hoặc số lượng bổ sung hàng.",
    "StockCountSession": "Bảng StockCountSession lưu một phiên kiểm kê, gồm phạm vi, mục đích, người thực hiện, ca làm việc, trạng thái, mốc thời gian và kết quả tổng hợp chênh lệch.",
    "StockCountItem": "Bảng StockCountItem lưu kết quả kiểm đếm từng sản phẩm trong một phiên kiểm kê, so sánh số lượng hệ thống với số lượng thực tế và ghi nhận chênh lệch, kiểm đếm lại, bằng chứng và biến động tồn.",
    "InvoiceIssueJob": "Bảng InvoiceIssueJob lưu một tác vụ phát hành hóa đơn điện tử cho một cơ sở, theo dõi phạm vi ngày kinh doanh, hạn phát hành, tiến độ và kết quả tổng hợp.",
    "InvoiceIssueJobItem": "Bảng InvoiceIssueJobItem lưu kết quả phát hành hóa đơn cho từng đơn hàng trong một tác vụ, gồm dữ liệu chuẩn bị, số lần thử, mã giao dịch, số hóa đơn và lỗi từ nhà cung cấp.",
    "InvoiceBulkIssueRun": "Bảng InvoiceBulkIssueRun lưu một lần người dùng chọn và khởi tạo phát hành hóa đơn hàng loạt, gồm tiêu chí chọn, danh sách đơn đủ điều kiện hoặc bị loại và các tác vụ phát hành được tạo.",
    "MeInvoiceAccount": "Bảng MeInvoiceAccount lưu cấu hình tài khoản kết nối dịch vụ MISA meInvoice của một pháp nhân, gồm môi trường, địa chỉ dịch vụ, trạng thái thông tin xác thực và kết quả kiểm tra kết nối.",
    "MeInvoiceStoreConfig": "Bảng MeInvoiceStoreConfig lưu cấu hình phát hành hóa đơn của từng cơ sở trên MISA meInvoice, gồm ký hiệu hóa đơn, cách tính thuế, quy tắc ánh xạ dữ liệu bán hàng và thông tin người mua mặc định.",
    "InvoiceSourceOrder": "Bảng InvoiceSourceOrder lưu đơn hàng nhận từ hệ thống bán hàng nguồn sau khi chuẩn hóa, đối soát số tiền và kiểm tra điều kiện lập hóa đơn; bản ghi này là nguồn để tạo hoặc cập nhật chứng từ hóa đơn.",
    "InvoiceDocument": "Bảng InvoiceDocument lưu chứng từ hóa đơn đã chuẩn bị từ đơn hàng nguồn, gồm thông tin người mua, hàng hóa, số tiền, kết quả kiểm tra, lịch sử chỉnh sửa/xét duyệt và trạng thái phát hành.",
    "InvoiceOrderSyncRun": "Bảng InvoiceOrderSyncRun ghi nhận một lần đồng bộ đơn hàng phục vụ hóa đơn theo cơ sở và ngày kinh doanh, gồm số đơn được thêm mới, cập nhật, giữ nguyên và lỗi xử lý.",
    "LeavePolicy": "Bảng LeavePolicy lưu chính sách phép năm, gồm thời điểm cộng phép hằng tháng, hạn mức năm, thời hạn dùng phép chuyển tiếp và quy tắc khóa phép trong thời gian thử việc.",
    "CompanyHoliday": "Bảng CompanyHoliday lưu các ngày nghỉ chung của công ty để loại trừ hoặc áp dụng đúng khi tính ngày làm việc và ngày nghỉ phép.",
    "LeaveRequest": "Bảng LeaveRequest lưu đơn nghỉ phép của nhân viên, gồm loại nghỉ, các ngày/buổi nghỉ, tổng số đơn vị phép, phân bổ số dư, trạng thái phê duyệt và nguồn tạo đơn.",
    "LeaveBalanceBucket": "Bảng LeaveBalanceBucket lưu số dư phép của nhân viên theo năm, tách số khả dụng, đang giữ cho đơn chờ duyệt, đã sử dụng, chờ hết thử việc và đã hết hạn.",
    "LeaveLedgerEntry": "Bảng LeaveLedgerEntry lưu từng bút toán làm tăng hoặc giảm số dư phép, giúp truy vết nguồn thay đổi từ đơn nghỉ, lô nhập hoặc nghiệp vụ điều chỉnh.",
    "LeaveDayReservation": "Bảng LeaveDayReservation đánh dấu buổi sáng và buổi chiều của một ngày đã được đơn nghỉ nào giữ, nhằm ngăn nhân viên tạo các đơn nghỉ trùng thời gian.",
    "LeaveImportRow": "Bảng LeaveImportRow lưu kết quả kiểm tra và ghi nhận từng dòng trong lô nhập dữ liệu nghỉ phép hoặc số dư phép.",
    "LeaveApprovalConfig": "Bảng LeaveApprovalConfig lưu cấu hình các cấp duyệt nghỉ phép theo phạm vi áp dụng, gồm thứ tự cấp duyệt và cách xác định người duyệt.",
    "LeaveApprovalTask": "Bảng LeaveApprovalTask lưu nhiệm vụ duyệt tại từng cấp của một đơn nghỉ phép, gồm người/nhóm được giao, trạng thái, quyết định và người thực hiện.",
    "LeaveApprovalReassignment": "Bảng LeaveApprovalReassignment lưu lịch sử chuyển giao nhiệm vụ duyệt nghỉ phép từ người hoặc nhóm cũ sang người hoặc nhóm mới.",
    "MarketingVoucherCampaign": "Bảng MarketingVoucherCampaign lưu cấu hình một chiến dịch voucher marketing, gồm loại và giá trị ưu đãi, thời hạn, quy tắc sinh mã, hình ảnh, trạng thái và số lượng mã.",
    "MarketingVoucherCode": "Bảng MarketingVoucherCode lưu từng mã voucher thuộc chiến dịch, theo dõi người nhận, thời điểm phân phối, sử dụng, gửi email hoặc thu hồi.",
    "MarketingVoucherJob": "Bảng MarketingVoucherJob lưu tác vụ nền dùng để sinh, gia hạn hoặc xuất voucher, gồm tiến độ, vị trí tệp kết quả, số lần thử và lỗi gần nhất.",
    "MarketingVoucherJobItem": "Bảng MarketingVoucherJobItem lưu phần việc cụ thể của một tác vụ voucher, gồm danh sách mã, người nhận email và kết quả gửi/xử lý.",
    "Organization": "Bảng Organization lưu thông tin pháp nhân hoặc tổ chức sở hữu các cơ sở trong hệ thống, gồm mã, tên, mã số thuế, địa chỉ và hình ảnh nhận diện.",
    "Warehouse": "Bảng Warehouse lưu thông tin cơ sở trong hệ thống, gồm đơn vị sở hữu, mã và tên cơ sở, loại hình, địa chỉ, người quản lý, trạng thái và tọa độ.",
    "WarehouseLocation": "Bảng WarehouseLocation lưu các khu vực lưu trữ hoặc khu vực nghiệp vụ bên trong một cơ sở, phục vụ quản lý vị trí hàng hóa và chứng từ kho.",
    "WarehouseLocationSlot": "Bảng WarehouseLocationSlot lưu ô/vị trí chi tiết bên trong một khu vực của cơ sở, gồm mã, thứ tự hiển thị và trạng thái sử dụng.",
    "WarehouseLocationSlotProduct": "Bảng WarehouseLocationSlotProduct lưu quan hệ giữa sản phẩm và ô/vị trí trưng bày hoặc lưu trữ, dùng để xác định sản phẩm xuất hiện ở vị trí nào và theo thứ tự nào.",
    "ProductCategory": "Bảng ProductCategory lưu cây danh mục sản phẩm, cho phép phân nhóm nhiều cấp để phục vụ tra cứu, hiển thị và báo cáo.",
    "Product": "Bảng Product lưu dữ liệu danh mục sản phẩm, gồm mã, mã vạch, phân loại, đặc tính kỹ thuật, xuất xứ, đơn vị tính, giá và yêu cầu quản lý số sê-ri.",
    "ProductBOM": "Bảng ProductBOM lưu cấu trúc định mức sản phẩm, xác định một sản phẩm cha cần bao nhiêu đơn vị của từng sản phẩm thành phần.",
    "InAppNotification": "Bảng InAppNotification lưu thông báo hiển thị trong ứng dụng cho người dùng hoặc vai trò, gồm nội dung, mức ưu tiên, liên kết thao tác và trạng thái đã đọc.",
    "NotificationPushToken": "Bảng NotificationPushToken lưu mã thiết bị dùng để gửi thông báo đẩy cho người dùng, cùng nền tảng, quyền nhận thông báo và trạng thái hoạt động của mã.",
    "PosDevice": "Bảng PosDevice lưu thiết bị bán hàng đã đăng ký tại một cơ sở, gồm dấu vân tay thiết bị, thông tin xác thực, phiên bản ứng dụng và trạng thái thu hồi.",
    "PosDeviceEnrollment": "Bảng PosDeviceEnrollment lưu mã đăng ký dùng một lần để liên kết thiết bị POS với cơ sở, theo dõi thời hạn, thiết bị đã sử dụng mã và trạng thái thu hồi.",
    "PosCustomerDisplaySettings": "Bảng PosCustomerDisplaySettings lưu cấu hình màn hình hiển thị cho khách tại một cơ sở, chủ yếu là phiên bản cấu hình và danh sách nội dung phát.",
    "PosCustomerDisplayMedia": "Bảng PosCustomerDisplayMedia lưu siêu dữ liệu của ảnh hoặc video dùng trên màn hình khách hàng POS, gồm vị trí tệp, kích thước, thời lượng và mã kiểm tra nội dung.",
    "PosProductVisibilitySettings": "Bảng PosProductVisibilitySettings lưu cấu hình ẩn nhóm hàng hoặc sản phẩm trên giao diện POS của một cơ sở.",
    "PosReceiptSettings": "Bảng PosReceiptSettings lưu cấu hình in hóa đơn bán lẻ của cơ sở, gồm khổ giấy, thông tin cửa hàng, logo, cỡ chữ, nội dung cuối hóa đơn và các tùy chọn hiển thị.",
    "PosTicketSettings": "Bảng PosTicketSettings lưu cấu hình in phiếu/số thứ tự tại POS, gồm kích thước giấy, nội dung, logo, mã QR, cỡ chữ và điều kiện in tự động.",
    "PosLuckyDrawSettings": "Bảng PosLuckyDrawSettings lưu cấu hình chương trình phiếu rút thăm tại POS của một cơ sở, gồm trạng thái áp dụng, nội dung in và số phiếu theo từng gói hàng.",
    "PosPaymentSettings": "Bảng PosPaymentSettings lưu cấu hình thanh toán chuyển khoản của thiết bị POS, gồm tài khoản nhận tiền và quy tắc chỉ chấp nhận số tiền cố định.",
    "PosOrderSummary": "Bảng PosOrderSummary lưu bản tóm tắt đơn hàng tại POS để tra cứu và đồng bộ, gồm trạng thái thanh toán, trạng thái đồng bộ, khách hàng, nhân viên, hàng hóa và tổng tiền.",
    "PosMemberCompensation": "Bảng PosMemberCompensation lưu giao dịch điều chỉnh hoặc bồi hoàn giá trị thành viên phát sinh tại POS, đồng thời theo dõi kết quả gửi sang hệ thống bên ngoài.",
    "ProcessConfig": "Bảng ProcessConfig lưu cấu hình quy trình phê duyệt theo loại đối tượng và cơ sở, gồm chuỗi người duyệt, điều kiện tự duyệt, yêu cầu bằng chứng và xác thực OTP.",
    "NonconformityReport": "Bảng NonconformityReport lưu báo cáo sai lệch hoặc không phù hợp phát hiện trong nhập, xuất, kiểm kê hay vận hành, gồm số lượng ảnh hưởng, bằng chứng, người xử lý và phương án giải quyết.",
    "QuarantineRecord": "Bảng QuarantineRecord lưu việc đưa hàng không phù hợp vào trạng thái cách ly và theo dõi số lượng, vị trí, lý do, người giải phóng và kết quả xử lý.",
    "ReportTemplateVersion": "Bảng ReportTemplateVersion lưu từng phiên bản tệp mẫu báo cáo, gồm vị trí lưu trữ, cấu trúc workbook, quy tắc ánh xạ dữ liệu và trạng thái sử dụng.",
    "RevenueOrderItem": "Bảng RevenueOrderItem biểu diễn dữ liệu đơn hàng đã tổng hợp cho báo cáo doanh thu, gồm trạng thái, thời gian, nhân viên, phương thức thanh toán, số lượng và các thành phần giá trị tiền.",
    "SoldOrderGoodsItem": "Bảng SoldOrderGoodsItem biểu diễn từng mặt hàng đã bán trong dữ liệu báo cáo, kết hợp thông tin đơn hàng, sản phẩm, số lượng, tiền bán, tiền giảm, tiền hủy và thuế.",
    "PosTransactionItem": "Bảng PosTransactionItem lưu từng dòng hàng trong một giao dịch POS, gồm sản phẩm, số lượng, đơn giá và thành tiền.",
    "AuditLog": "Bảng AuditLog lưu nhật ký thay đổi dữ liệu và thao tác người dùng, gồm đối tượng bị tác động, hành động, giá trị trước/sau, thiết bị, địa chỉ IP và thời điểm đồng bộ.",
    "ExternalStoreBinding": "Bảng ExternalStoreBinding ánh xạ một cửa hàng/tài khoản trên hệ thống bên ngoài với một hoặc nhiều cơ sở trong hệ thống, phục vụ đồng bộ dữ liệu đúng phạm vi.",
    "User": "Bảng User lưu tài khoản đăng nhập và trạng thái truy cập của người dùng, cùng thông tin nhân viên, cơ sở làm việc, xác thực đa yếu tố và dữ liệu OTP email.",
    "Role": "Bảng Role lưu vai trò và tập quyền trong hệ thống, hỗ trợ quan hệ vai trò cha-con, màu hiển thị và thứ tự trên bảng công việc.",
    "UserWarehouseRole": "Bảng UserWarehouseRole gán vai trò cho người dùng tại một cơ sở trong khoảng thời gian xác định; tên bảng giữ thuật ngữ kỹ thuật Warehouse nhưng nghiệp vụ sử dụng khái niệm cơ sở.",
    "ImportVoucher": "Bảng ImportVoucher lưu phiếu nhập hàng của một cơ sở, gồm nhà cung cấp, đơn mua tham chiếu, trạng thái phê duyệt, người lập/duyệt và chứng từ đính kèm.",
    "ImportVoucherItem": "Bảng ImportVoucherItem lưu từng mặt hàng của phiếu nhập, gồm sản phẩm, khu vực nhận hàng, số lượng dự kiến/thực nhận, đơn giá và tình trạng hàng.",
    "ExportVoucher": "Bảng ExportVoucher lưu phiếu xuất hàng của một cơ sở, gồm mục đích xuất, đối tượng nhận, chứng từ tham chiếu, trạng thái phê duyệt và việc trừ tồn khả dụng.",
    "ExportVoucherItem": "Bảng ExportVoucherItem lưu từng mặt hàng của phiếu xuất, gồm sản phẩm, khu vực lấy hàng, số lượng cần xuất, số lượng đã lấy và đơn giá.",
    "TransferOrder": "Bảng TransferOrder lưu lệnh điều chuyển hàng giữa hai cơ sở, theo dõi phê duyệt, xuất hàng, nhận hàng, xác thực lại và các chứng từ liên quan.",
    "TransferOrderItem": "Bảng TransferOrderItem lưu từng sản phẩm trong lệnh điều chuyển, gồm vị trí nguồn/đích, số lượng yêu cầu, số lượng đã nhận và trạng thái dòng hàng.",
    "WorkflowDefinition": "Bảng WorkflowDefinition lưu định nghĩa tổng quát của một quy trình, gồm loại đối tượng áp dụng, phạm vi cơ sở, trạng thái và phiên bản đang dùng.",
    "WorkflowVersion": "Bảng WorkflowVersion lưu một phiên bản bất biến của định nghĩa quy trình dưới dạng các nút và liên kết, cùng thông tin công bố.",
    "WorkflowInstance": "Bảng WorkflowInstance lưu một lần thực thi quy trình cho một đối tượng nghiệp vụ cụ thể, theo dõi phiên bản, các bước hiện tại, người khởi tạo và trạng thái hoàn tất.",
    "WorkflowTask": "Bảng WorkflowTask lưu nhiệm vụ phát sinh tại một nút của phiên bản quy trình, gồm người/vai trò được giao, kết quả, hạn xử lý và người hoàn thành.",
}


ENTITY_LABELS = {
    "WarehouseAttendancePolicy": "chính sách chấm công của cơ sở",
    "AttendanceWorkArrangement": "thỏa thuận làm việc ngoài cơ sở",
    "WarehouseAttendanceExemption": "ngoại lệ chấm công",
    "AttendanceLog": "lượt chấm công",
    "AttendanceLateReport": "báo cáo đi muộn",
    "EmployeeContract": "hợp đồng lao động",
    "EmployeeContractAutomationRun": "lần chạy tự động xử lý hợp đồng",
    "EmployeeContractDocument": "tài liệu hợp đồng",
    "EmployeeContractDocumentUploadIntent": "yêu cầu tải tài liệu hợp đồng",
    "EmployeeContractImportRow": "dòng dữ liệu nhập hợp đồng",
    "EmployeeProfile": "hồ sơ nhân viên",
    "EmployeeEmploymentTransition": "yêu cầu chuyển trạng thái làm việc",
    "ExternalScanQueue": "lượt quét hàng bên ngoài",
    "ExternalQueueScannableProductConfig": "cấu hình sản phẩm được phép quét",
    "OfficeScopeConfig": "cấu hình phạm vi cơ sở",
    "OfficeScopeEdge": "quan hệ phạm vi cơ sở",
    "OfficeScopeMaterialization": "lần tính phạm vi truy cập",
    "UserFacilityAccessGrant": "quyền truy cập cơ sở của người dùng",
    "UserAccessVersion": "phiên bản quyền truy cập người dùng",
    "UserAccessMetadata": "thông tin tổng hợp quyền truy cập",
    "FileTemplate": "tệp biểu mẫu",
    "Inventory": "bản ghi tồn kho",
    "InventoryStockPolicy": "chính sách ngưỡng tồn",
    "StockCountSession": "phiên kiểm kê",
    "StockCountItem": "dòng kiểm kê",
    "InvoiceIssueJob": "tác vụ phát hành hóa đơn",
    "InvoiceIssueJobItem": "dòng xử lý phát hành hóa đơn",
    "InvoiceBulkIssueRun": "lần phát hành hóa đơn hàng loạt",
    "MeInvoiceAccount": "tài khoản kết nối meInvoice",
    "MeInvoiceStoreConfig": "cấu hình meInvoice của cơ sở",
    "InvoiceSourceOrder": "đơn hàng nguồn lập hóa đơn",
    "InvoiceDocument": "chứng từ hóa đơn",
    "InvoiceOrderSyncRun": "lần đồng bộ đơn hàng hóa đơn",
    "LeavePolicy": "chính sách nghỉ phép",
    "CompanyHoliday": "ngày nghỉ công ty",
    "LeaveRequest": "đơn nghỉ phép",
    "LeaveBalanceBucket": "số dư phép theo năm",
    "LeaveLedgerEntry": "bút toán phép",
    "LeaveDayReservation": "phần ngày nghỉ đã giữ",
    "LeaveImportRow": "dòng dữ liệu nhập nghỉ phép",
    "LeaveApprovalConfig": "cấu hình duyệt nghỉ phép",
    "LeaveApprovalTask": "nhiệm vụ duyệt nghỉ phép",
    "LeaveApprovalReassignment": "lần chuyển giao nhiệm vụ duyệt",
    "MarketingVoucherCampaign": "chiến dịch voucher",
    "MarketingVoucherCode": "mã voucher",
    "MarketingVoucherJob": "tác vụ xử lý voucher",
    "MarketingVoucherJobItem": "dòng xử lý voucher",
    "Organization": "tổ chức",
    "Warehouse": "cơ sở",
    "WarehouseLocation": "khu vực trong cơ sở",
    "WarehouseLocationSlot": "ô/vị trí trong khu vực",
    "WarehouseLocationSlotProduct": "quan hệ sản phẩm với ô/vị trí",
    "ProductCategory": "danh mục sản phẩm",
    "Product": "sản phẩm",
    "ProductBOM": "dòng định mức sản phẩm",
    "InAppNotification": "thông báo trong ứng dụng",
    "NotificationPushToken": "mã nhận thông báo đẩy",
    "PosDevice": "thiết bị POS",
    "PosDeviceEnrollment": "mã đăng ký thiết bị POS",
    "PosCustomerDisplaySettings": "cấu hình màn hình khách hàng POS",
    "PosCustomerDisplayMedia": "tệp media màn hình khách hàng",
    "PosProductVisibilitySettings": "cấu hình hiển thị sản phẩm POS",
    "PosReceiptSettings": "cấu hình hóa đơn in tại POS",
    "PosTicketSettings": "cấu hình phiếu in tại POS",
    "PosLuckyDrawSettings": "cấu hình phiếu rút thăm POS",
    "PosPaymentSettings": "cấu hình thanh toán POS",
    "PosOrderSummary": "bản tóm tắt đơn POS",
    "PosMemberCompensation": "giao dịch điều chỉnh thành viên",
    "ProcessConfig": "cấu hình quy trình phê duyệt",
    "NonconformityReport": "báo cáo không phù hợp",
    "QuarantineRecord": "bản ghi cách ly hàng",
    "ReportTemplateVersion": "phiên bản mẫu báo cáo",
    "RevenueOrderItem": "dòng dữ liệu đơn hàng doanh thu",
    "SoldOrderGoodsItem": "dòng hàng đã bán",
    "PosTransactionItem": "dòng hàng giao dịch POS",
    "AuditLog": "nhật ký kiểm toán",
    "ExternalStoreBinding": "ánh xạ cửa hàng bên ngoài",
    "User": "tài khoản người dùng",
    "Role": "vai trò",
    "UserWarehouseRole": "vai trò người dùng tại cơ sở",
    "ImportVoucher": "phiếu nhập",
    "ImportVoucherItem": "dòng hàng phiếu nhập",
    "ExportVoucher": "phiếu xuất",
    "ExportVoucherItem": "dòng hàng phiếu xuất",
    "TransferOrder": "lệnh điều chuyển",
    "TransferOrderItem": "dòng hàng điều chuyển",
    "WorkflowDefinition": "định nghĩa quy trình",
    "WorkflowVersion": "phiên bản quy trình",
    "WorkflowInstance": "lần thực thi quy trình",
    "WorkflowTask": "nhiệm vụ quy trình",
}


REFERENCE_LABELS = {
    "warehouse": "cơ sở",
    "workplace_warehouse": "cơ sở làm việc",
    "source_warehouse": "cơ sở xuất",
    "destination_warehouse": "cơ sở nhận",
    "warehouse_location": "khu vực trong cơ sở",
    "warehouse_location_slot": "ô/vị trí trong khu vực",
    "source_location": "khu vực xuất",
    "destination_location": "khu vực nhận",
    "organization": "tổ chức",
    "user": "người dùng",
    "target_user": "người dùng nhận thông báo",
    "employee_user": "tài khoản người dùng của nhân viên",
    "employee_profile": "hồ sơ nhân viên",
    "employee": "nhân viên",
    "product": "sản phẩm",
    "parent_product": "sản phẩm cha",
    "child_product": "sản phẩm thành phần",
    "category": "danh mục sản phẩm",
    "parent": "bản ghi cha",
    "role": "vai trò",
    "target_role": "vai trò nhận thông báo",
    "assigned_role": "vai trò được giao xử lý",
    "policy": "chính sách chấm công",
    "work_arrangement": "thỏa thuận làm việc ngoài cơ sở",
    "source_leave_request": "đơn nghỉ phép nguồn",
    "attendance_log": "bản ghi chấm công",
    "contract": "hợp đồng lao động",
    "renewed_from_contract": "hợp đồng được gia hạn",
    "root_contract": "hợp đồng gốc trong chuỗi gia hạn",
    "upload_intent": "yêu cầu tải tệp",
    "finalized_document": "tài liệu đã hoàn tất",
    "batch": "lô nhập dữ liệu",
    "document": "tài liệu được tạo",
    "inventory": "bản ghi tồn kho",
    "session": "phiên xử lý",
    "previous_access_session": "phiên truy cập trước",
    "handover_to_session": "phiên nhận bàn giao",
    "access_session": "phiên truy cập",
    "external_client": "ứng dụng khách bên ngoài",
    "device": "thiết bị",
    "used_by_device": "thiết bị đã dùng mã đăng ký",
    "job": "tác vụ",
    "invoice_document": "chứng từ hóa đơn",
    "source_order": "đơn hàng nguồn",
    "source_order_document": "đơn hàng nguồn đã chuẩn hóa",
    "last_sync_run": "lần đồng bộ gần nhất",
    "meinvoice_account": "tài khoản meInvoice",
    "legal_entity": "pháp nhân phát hành hóa đơn",
    "leave_request": "đơn nghỉ phép",
    "request": "yêu cầu nghiệp vụ",
    "import_batch": "lô nhập dữ liệu",
    "last_ledger_entry": "bút toán phép gần nhất",
    "morning_request": "đơn nghỉ giữ buổi sáng",
    "afternoon_request": "đơn nghỉ giữ buổi chiều",
    "approval_task": "nhiệm vụ duyệt",
    "campaign": "chiến dịch voucher",
    "active_generation_job": "tác vụ sinh mã đang hoạt động",
    "active_extension_job": "tác vụ gia hạn đang hoạt động",
    "source_instance": "lần thực thi quy trình nguồn",
    "source_entity": "đối tượng nghiệp vụ nguồn",
    "transaction": "giao dịch POS",
    "local_order": "đơn hàng cục bộ trên POS",
    "remote_order": "đơn hàng trên hệ thống bên ngoài",
    "shop": "cửa hàng trên hệ thống bên ngoài",
    "member": "thành viên trên hệ thống bên ngoài",
    "nonconformity_report": "báo cáo không phù hợp",
    "source_item": "dòng chứng từ nguồn",
    "reporter": "người lập báo cáo",
    "reviewer": "người xem xét",
    "resolved_by": "người xử lý hoàn tất",
    "template": "mẫu báo cáo",
    "entity": "đối tượng nghiệp vụ",
    "canonical_warehouse": "cơ sở chính được ánh xạ",
    "current_version": "phiên bản quy trình đang dùng",
    "workflow_definition": "định nghĩa quy trình",
    "workflow_version": "phiên bản quy trình",
    "instance": "lần thực thi quy trình",
    "node": "nút trong quy trình",
    "export_voucher": "phiếu xuất",
    "import_voucher": "phiếu nhập",
    "transfer_order": "lệnh điều chuyển",
    "purchase_order": "đơn mua hàng",
    "reference": "chứng từ hoặc đối tượng tham chiếu",
    "office": "văn phòng/cơ sở",
    "target_facility": "cơ sở đích",
    "facility": "cơ sở được cấp quyền",
    "workplace_facility": "cơ sở làm việc",
    "access_version": "phiên bản quyền truy cập",
}


EXACT_DESCRIPTIONS = {
    "id": "Mã định danh duy nhất của {entity}.",
    "warehouse_id": "Mã cơ sở nơi {entity} phát sinh hoặc được áp dụng.",
    "created_at": "Thời điểm bản ghi được tạo trên hệ thống.",
    "updated_at": "Thời điểm nội dung bản ghi được cập nhật gần nhất.",
    "created_by": "Mã người dùng tạo bản ghi.",
    "updated_by": "Mã người dùng cập nhật bản ghi gần nhất.",
    "is_deleted": "Cờ xóa mềm: true khi bản ghi đã bị loại khỏi sử dụng nhưng vẫn được lưu để truy vết; false khi bản ghi còn hiệu lực trong dữ liệu.",
    "enabled": "Cho biết cấu hình hoặc chức năng có đang được bật để áp dụng hay không.",
    "is_active": "Cho biết bản ghi có đang được sử dụng để xử lý nghiệp vụ hay không.",
    "status": "Trạng thái hiện tại của {entity}, dùng để kiểm soát bước xử lý tiếp theo.",
    "notes": "Ghi chú bổ sung phục vụ xử lý và tra cứu {entity}.",
    "note": "Ghi chú bổ sung cho {entity}.",
    "description": "Nội dung mô tả giúp người dùng nhận biết mục đích hoặc phạm vi của {entity}.",
    "name": "Tên hiển thị của {entity}.",
    "code": "Mã nghiệp vụ dùng để nhận biết và tra cứu {entity}.",
    "version": "Số phiên bản của {entity}, tăng khi cấu hình hoặc nội dung được thay đổi.",
    "revision": "Số lần hiệu chỉnh của bản ghi, dùng để kiểm soát cập nhật đồng thời và truy vết thay đổi.",
    "action_time": "Thời điểm thao tác thực tế xảy ra trên thiết bị, kể cả khi thiết bị đang ngoại tuyến.",
    "sync_time": "Thời điểm dữ liệu được đồng bộ thành công lên máy chủ.",
    "effective_from": "Thời điểm bắt đầu áp dụng bản ghi.",
    "effective_to": "Thời điểm kết thúc áp dụng bản ghi; để trống khi chưa xác định ngày hết hiệu lực.",
    "valid_from": "Thời điểm bắt đầu có hiệu lực của bản ghi.",
    "valid_until": "Thời điểm hết hiệu lực của bản ghi; để trống nếu không giới hạn thời gian.",
    "start_date": "Ngày bắt đầu khoảng thời gian áp dụng.",
    "end_date": "Ngày kết thúc khoảng thời gian áp dụng.",
    "reason": "Lý do phát sinh hoặc thực hiện {entity}.",
    "cancellation_reason": "Lý do hủy {entity}.",
    "cancel_reason": "Lý do hủy {entity}.",
    "cancelled_by": "Mã người dùng thực hiện hủy {entity}.",
    "cancelled_at": "Thời điểm {entity} bị hủy.",
    "approved_by": "Mã người dùng phê duyệt {entity}.",
    "approved_at": "Thời điểm {entity} được phê duyệt.",
    "requested_by": "Mã người dùng khởi tạo yêu cầu hoặc tác vụ.",
    "completed_at": "Thời điểm {entity} hoàn tất xử lý.",
    "started_at": "Thời điểm bắt đầu xử lý {entity}.",
    "failed_at": "Thời điểm tác vụ được xác định là thất bại.",
    "error_message": "Thông báo chi tiết của lỗi làm tác vụ thất bại.",
    "last_error": "Nội dung lỗi gần nhất của lần xử lý {entity}.",
    "last_error_code": "Mã lỗi gần nhất, dùng để phân loại nguyên nhân và quyết định có thể thử lại hay không.",
    "error_code": "Mã lỗi chuẩn hóa phát sinh khi xử lý bản ghi.",
    "attempt_count": "Tổng số lần hệ thống đã thử xử lý {entity}.",
    "attempt": "Số thứ tự lần thử xử lý hiện tại.",
    "idempotency_key": "Khóa chống xử lý trùng; các yêu cầu lặp lại có cùng khóa không được tạo kết quả nghiệp vụ mới.",
    "timezone": "Múi giờ dùng để quy đổi ngày và thời điểm nghiệp vụ; dữ liệu hiện áp dụng múi giờ Asia/Ho_Chi_Minh.",
    "employee_id": "Mã nhân viên nghiệp vụ tại thời điểm ghi nhận, được lưu kèm để tra cứu và đối soát.",
    "employee_name": "Tên nhân viên tại thời điểm ghi nhận, được lưu kèm để hiển thị và giữ nguyên lịch sử.",
    "employee_code": "Mã nhân viên dùng để đối chiếu hồ sơ nhân sự khi nhập hoặc xử lý dữ liệu.",
    "full_name": "Họ và tên đầy đủ của nhân viên.",
    "email": "Địa chỉ email dùng để liên hệ hoặc đăng nhập, tùy nghiệp vụ của tài khoản.",
    "phone": "Số điện thoại liên hệ của nhân viên.",
    "job_title": "Chức danh công việc hiện tại của nhân viên.",
    "department": "Phòng ban hoặc bộ phận mà nhân viên đang làm việc.",
    "ip_addresses": "Danh sách địa chỉ IP được coi là mạng hợp lệ của cơ sở khi xác minh chấm công.",
    "verification_strategy": "Chiến lược xác minh chấm công mà cơ sở áp dụng, ví dụ kiểm tra IP, GPS hoặc kết hợp nhiều điều kiện.",
    "gps_radius_m": "Bán kính cho phép chấm công quanh tọa độ mục tiêu của cơ sở, tính bằng mét.",
    "gps_max_accuracy_m": "Sai số GPS lớn nhất được chấp nhận, tính bằng mét; dữ liệu vị trí có độ chính xác kém hơn ngưỡng này bị từ chối.",
    "gps_max_age_seconds": "Tuổi tối đa của dữ liệu GPS tại thời điểm chấm công, tính bằng giây; dữ liệu cũ hơn ngưỡng này không được chấp nhận.",
    "allow_business_trip": "Cho biết chính sách cơ sở có cho phép chấm công theo thỏa thuận đi công tác đã được duyệt hay không.",
    "allow_work_from_home": "Cho biết chính sách cơ sở có cho phép chấm công theo thỏa thuận làm việc tại nhà đã được duyệt hay không.",
    "attendance_required": "Cho biết người dùng có bắt buộc chấm công tại cơ sở trong khoảng thời gian hiệu lực hay không.",
    "attendance_date": "Ngày làm việc mà lượt chấm công hoặc báo cáo đi muộn được ghi nhận.",
    "check_in_at": "Thời điểm chấm công được hệ thống ghi nhận sau khi xác minh.",
    "check_in_method": "Phương thức nhân viên dùng để chấm công, ví dụ qua mạng nội bộ hoặc vị trí GPS.",
    "work_mode": "Chế độ làm việc tại thời điểm chấm công, như làm tại cơ sở, làm tại nhà hoặc đi công tác.",
    "location": "Tọa độ và thông tin độ chính xác do thiết bị gửi khi chấm công.",
    "distance_from_target_m": "Khoảng cách từ vị trí thiết bị đến tọa độ mục tiêu được áp dụng, tính bằng mét.",
    "rejected_reason": "Mã lý do hệ thống từ chối lượt chấm công, ví dụ sai IP, ngoài bán kính hoặc dữ liệu GPS không đạt yêu cầu.",
    "location_rule": "Quy tắc vị trí áp dụng cho thỏa thuận làm việc, xác định có cần kiểm tra tọa độ và bán kính hay không.",
    "destination_name": "Tên địa điểm công tác hoặc nơi làm việc ngoài cơ sở để người duyệt và nhân viên nhận biết.",
    "destination_coordinate": "Tọa độ địa điểm làm việc ngoài cơ sở dùng làm tâm kiểm tra vị trí khi chấm công.",
    "radius_m": "Bán kính chấm công cho phép quanh địa điểm của thỏa thuận làm việc, tính bằng mét.",
    "expected_arrival_time": "Giờ nhân viên được yêu cầu có mặt theo lịch làm việc trong ngày.",
    "estimated_arrival_time": "Giờ nhân viên dự kiến đến cơ sở theo báo cáo đi muộn.",
    "reviewed_by": "Mã người dùng đã xem xét báo cáo hoặc chứng từ.",
    "reviewed_at": "Thời điểm báo cáo hoặc chứng từ được xem xét.",
    "review_notes": "Nhận xét của người xem xét, dùng để giải thích quyết định chấp nhận hoặc từ chối.",
    "contract_number": "Số hợp đồng lao động hiển thị trên hồ sơ và tài liệu chính thức.",
    "contract_number_normalized": "Số hợp đồng đã được chuẩn hóa để tìm kiếm, so sánh và ngăn tạo trùng.",
    "contract_type": "Loại hợp đồng lao động, dùng để áp dụng đúng quy tắc thời hạn và gia hạn.",
    "renewal_sequence": "Số thứ tự của hợp đồng trong chuỗi gia hạn, bắt đầu từ hợp đồng gốc.",
    "termination_date": "Ngày chấm dứt hiệu lực hợp đồng lao động.",
    "termination_reason": "Lý do chấm dứt hợp đồng lao động.",
    "terminated_by": "Mã người dùng ghi nhận việc chấm dứt hợp đồng.",
    "terminated_at": "Thời điểm thao tác chấm dứt hợp đồng được ghi nhận trên hệ thống.",
    "job_type": "Loại tác vụ tự động liên quan đến hợp đồng được thực hiện trong lần chạy.",
    "as_of_date": "Ngày mốc mà tác vụ dùng để xác định hợp đồng cần xử lý.",
    "lease_expires_at": "Thời điểm quyền giữ tác vụ của tiến trình hiện tại hết hạn; sau mốc này tiến trình khác có thể nhận xử lý lại.",
    "result": "Kết quả có cấu trúc do tác vụ hoặc bước quy trình tạo ra sau khi xử lý.",
    "storage_path": "Đường dẫn nội bộ của tệp trong dịch vụ lưu trữ.",
    "storage_generation": "Mã thế hệ của tệp trong dịch vụ lưu trữ, dùng để nhận biết đúng phiên bản vật lý của tệp.",
    "original_file_name": "Tên tệp gốc do người dùng cung cấp khi tải lên.",
    "mime_type": "Loại nội dung MIME của tệp, dùng để kiểm tra định dạng và phục vụ tải xuống.",
    "file_size": "Dung lượng tệp, tính bằng byte.",
    "file_size_bytes": "Dung lượng tệp, tính bằng byte.",
    "sha256": "Mã băm SHA-256 của nội dung tệp, dùng để kiểm tra toàn vẹn và phát hiện tệp trùng.",
    "checksum_sha256": "Mã băm SHA-256 của nội dung tệp, dùng để kiểm tra toàn vẹn và phát hiện thay đổi.",
    "is_current": "Cho biết đây có phải phiên bản tài liệu hợp đồng đang được sử dụng hay không.",
    "uploaded_by": "Mã người dùng tải tệp lên hệ thống.",
    "upload_storage_path": "Đường dẫn lưu trữ tạm/đích được cấp cho lần tải tệp hợp đồng.",
    "expected_mime_type": "Loại MIME mà tệp tải lên phải khớp để được chấp nhận.",
    "max_file_size": "Dung lượng tệp tối đa được phép tải lên, tính bằng byte.",
    "request_hash": "Mã băm của nội dung yêu cầu tải tệp, dùng để nhận biết yêu cầu lặp hoặc bị thay đổi.",
    "expires_at": "Thời điểm yêu cầu hoặc mã tạm hết hiệu lực và không còn được sử dụng.",
    "finalized_at": "Thời điểm quá trình tải tệp được xác nhận hoàn tất.",
    "failure_code": "Mã nguyên nhân khiến yêu cầu không thể hoàn tất.",
    "row_number": "Số thứ tự dòng trong tệp nhập nguồn, dùng để báo lỗi và đối soát.",
    "source_reference": "Mã tham chiếu từ dữ liệu nguồn, dùng để truy vết và tránh ghi nhận trùng.",
    "normalized_payload": "Dữ liệu của dòng nhập sau khi chuẩn hóa về cấu trúc và định dạng mà hệ thống sử dụng.",
    "validation_messages": "Danh sách lỗi hoặc cảnh báo phát hiện khi kiểm tra dữ liệu của dòng nhập.",
    "staged_document": "Thông tin tài liệu đã được chuẩn bị tạm thời để gắn với hợp đồng khi dòng nhập được ghi nhận.",
    "committed_at": "Thời điểm dòng nhập được ghi nhận thành công vào dữ liệu nghiệp vụ.",
    "probation_start_date": "Ngày nhân viên bắt đầu thời gian thử việc.",
    "probation_end_date": "Ngày dự kiến hoặc thực tế kết thúc thời gian thử việc.",
    "official_start_date": "Ngày nhân viên bắt đầu làm việc chính thức.",
    "resignation_date": "Ngày nghỉ việc của nhân viên.",
    "employment_status": "Trạng thái quan hệ lao động hiện tại của nhân viên, như thử việc, chính thức hoặc đã nghỉ việc.",
    "from_status": "Trạng thái làm việc của nhân viên trước khi chuyển đổi.",
    "to_status": "Trạng thái làm việc dự kiến sau khi chuyển đổi được áp dụng.",
    "effective_date": "Ngày thay đổi trạng thái hoặc chính sách bắt đầu có hiệu lực.",
    "applied_by": "Mã người dùng đã áp dụng thay đổi trạng thái nhân viên.",
    "applied_at": "Thời điểm thay đổi trạng thái nhân viên được áp dụng.",
    "client_id": "Mã ứng dụng khách bên ngoài đã gửi dữ liệu quét.",
    "barcode_scanned": "Giá trị mã vạch được thiết bị đọc tại thời điểm quét hàng.",
    "quantity": "Số lượng sản phẩm của dòng nghiệp vụ.",
    "unit_price": "Đơn giá của một đơn vị sản phẩm tại thời điểm ghi nhận.",
    "scan_time": "Thời điểm thiết bị thực hiện quét mã hàng.",
    "operator_name": "Tên người vận hành bên ngoài tại thời điểm quét.",
    "operator_id_external": "Mã người vận hành theo hệ thống bên ngoài.",
    "shift_id": "Mã ca làm việc gắn với hoạt động nghiệp vụ.",
    "shift_date": "Ngày kinh doanh của ca làm việc.",
    "batch_id": "Mã lô dùng để nhóm các bản ghi được nhập hoặc xử lý cùng một lần.",
    "final_approved_by": "Mã người dùng thực hiện bước phê duyệt cuối cùng.",
    "final_approved_at": "Thời điểm hoàn tất bước phê duyệt cuối cùng.",
    "revision_requested_by": "Mã người dùng yêu cầu chỉnh sửa bản ghi trước khi phê duyệt lại.",
    "revision_requested_at": "Thời điểm yêu cầu chỉnh sửa được tạo.",
    "rejection_reason": "Lý do bản ghi bị từ chối.",
    "atp_held": "Cho biết số lượng tương ứng đã được giữ khỏi tồn khả dụng (ATP) để tránh được sử dụng cho giao dịch khác hay chưa.",
    "product_ids": "Danh sách mã sản phẩm được cấu hình cho phép thực hiện nghiệp vụ.",
    "scope_mode": "Cách xác định phạm vi cơ sở mà văn phòng/cơ sở được phép truy cập.",
    "policy_version": "Phiên bản chính sách được dùng khi tính hoặc áp dụng quyền truy cập.",
    "scope_revision": "Số lần hiệu chỉnh cấu hình phạm vi được dùng cho lần tính này.",
    "requested_count": "Tổng số bản ghi dự kiến phải tính hoặc tạo trong lần xử lý.",
    "completed_count": "Số bản ghi đã xử lý thành công.",
    "failed_count": "Số bản ghi xử lý thất bại.",
    "attempts": "Tổng số lần hệ thống đã thử thực hiện tác vụ.",
    "permissions": "Danh sách quyền cụ thể mà người dùng hoặc vai trò được phép thực hiện.",
    "sources": "Danh sách nguồn tạo nên quyền truy cập, ví dụ nơi làm việc, vai trò được gán hoặc quyền quản trị.",
    "access_version": "Số phiên bản quyền truy cập đã được dùng để tạo bản ghi.",
    "computed_at": "Thời điểm hệ thống tính xong dữ liệu quyền truy cập.",
    "version_number": "Số thứ tự phiên bản, dùng để sắp xếp và xác định phiên bản mới hơn.",
    "source_fingerprint": "Dấu vân tay của các nguồn cấp quyền; thay đổi giá trị này cho biết quyền cần được tính lại.",
    "is_global_admin": "Cho biết người dùng có quyền quản trị trên toàn bộ cơ sở của hệ thống hay không.",
    "system_admin_sources": "Danh sách nguồn hoặc vai trò làm phát sinh quyền quản trị hệ thống.",
    "facility_grant_count": "Tổng số cơ sở mà người dùng được cấp quyền trong phiên bản này.",
    "computed_by": "Mã tác nhân hoặc tiến trình đã thực hiện tính quyền truy cập.",
    "activated_at": "Thời điểm phiên bản quyền truy cập được kích hoạt.",
    "retired_at": "Thời điểm phiên bản quyền truy cập ngừng được sử dụng.",
    "migration_version": "Phiên bản chuyển đổi dữ liệu đã tạo hoặc cập nhật bản ghi.",
    "title": "Tiêu đề hiển thị của {entity}.",
    "category": "Nhóm phân loại dùng để sắp xếp và lọc {entity}.",
    "file_name": "Tên tệp hiển thị cho người dùng.",
    "file_url": "Địa chỉ dùng để tải hoặc truy cập tệp.",
    "file_format": "Định dạng tệp, dùng để chọn cách mở hoặc xử lý phù hợp.",
    "version_history": "Danh sách các phiên bản trước của tệp và thông tin thay đổi tương ứng.",
    "total_quantity": "Tổng số lượng sản phẩm đang được ghi nhận tại vị trí tồn.",
    "atp_quantity": "Số lượng khả dụng để bán hoặc xuất (Available to Promise) sau khi trừ phần đã giữ và không sử dụng được.",
    "on_hold_quantity": "Số lượng đang được giữ cho giao dịch hoặc chờ xử lý nên chưa thể phân bổ tiếp.",
    "in_transit_quantity": "Số lượng đang trên đường điều chuyển đến hoặc đi khỏi cơ sở.",
    "quarantine_quantity": "Số lượng đang cách ly do vấn đề chất lượng và chưa được phép sử dụng.",
    "last_count_at": "Thời điểm sản phẩm tại vị trí này được kiểm kê gần nhất.",
    "last_updated_at": "Thời điểm số liệu tồn kho được cập nhật gần nhất.",
    "scope": "Phạm vi áp dụng của chính sách hoặc cấu hình, ví dụ toàn hệ thống, cơ sở, khu vực hoặc sản phẩm.",
    "min_stock_quantity": "Ngưỡng tồn tối thiểu; tồn thấp hơn mức này cần được cảnh báo hoặc bổ sung.",
    "max_stock_quantity": "Ngưỡng tồn tối đa cho phép tại phạm vi cấu hình.",
    "reorder_point_quantity": "Mức tồn kích hoạt đề xuất hoặc quy trình bổ sung hàng.",
    "reorder_quantity": "Số lượng đề xuất bổ sung khi tồn đạt hoặc thấp hơn điểm đặt hàng lại.",
    "session_number": "Số phiên kiểm kê hiển thị cho người dùng và dùng để tra cứu.",
    "count_scope": "Phạm vi hàng hóa cần kiểm trong phiên, ví dụ toàn cơ sở, khu vực hoặc nhóm sản phẩm.",
    "criteria": "Bộ tiêu chí chi tiết dùng để chọn các sản phẩm/vị trí đưa vào phiên kiểm kê.",
    "count_type": "Loại kiểm kê, xác định cách tổ chức và xử lý kết quả kiểm đếm.",
    "count_purpose": "Mục đích kiểm kê, dùng để giải thích lý do mở phiên và áp dụng quy trình phù hợp.",
    "checkpoint_type": "Loại điểm kiểm soát mà phiên kiểm kê được thực hiện.",
    "source": "Nguồn tạo bản ghi, giúp phân biệt dữ liệu do người dùng, quy trình tự động, tệp nhập hoặc hệ thống ngoài tạo ra.",
    "assigned_counter_ids": "Danh sách mã nhân viên được giao thực hiện kiểm đếm.",
    "counter_id": "Mã nhân viên trực tiếp chịu trách nhiệm kiểm đếm chính.",
    "supervisor_id": "Mã người giám sát phiên kiểm kê.",
    "external_operator_name": "Tên người kiểm đếm bên ngoài khi người thực hiện không có tài khoản nội bộ.",
    "external_operator_id": "Mã người kiểm đếm theo hệ thống bên ngoài.",
    "authorized_operator_ids": "Danh sách mã người được phép thao tác trong phiên kiểm kê.",
    "access_activated_at": "Thời điểm quyền truy cập phục vụ phiên kiểm kê được kích hoạt.",
    "finalized_batch_id": "Mã lô dữ liệu quét đã được chốt để dùng làm kết quả kiểm kê.",
    "business_date": "Ngày kinh doanh dùng để nhóm và đối soát giao dịch, có thể khác ngày dương lịch khi ca làm việc qua nửa đêm.",
    "blind_count_enabled": "Cho biết người kiểm đếm có bị ẩn số lượng tồn trên hệ thống để bảo đảm kiểm kê độc lập hay không.",
    "submitted_at": "Thời điểm kết quả được gửi để duyệt hoặc xử lý tiếp.",
    "discrepancy_count": "Tổng số dòng kiểm kê có chênh lệch giữa số hệ thống và số thực tế.",
    "system_quantity": "Số lượng tồn theo hệ thống tại mốc dữ liệu dùng để so sánh kiểm kê.",
    "atp_snapshot": "Số lượng tồn khả dụng được chụp lại khi tạo dòng kiểm kê.",
    "expected_at_count_time": "Số lượng hệ thống kỳ vọng tại đúng thời điểm kiểm đếm sau khi tính các biến động liên quan.",
    "current_atp": "Số lượng tồn khả dụng hiện tại khi dòng kiểm kê được xem hoặc xử lý.",
    "counted_quantity": "Số lượng thực tế mà người kiểm kê đã đếm được.",
    "counted_at": "Thời điểm hoàn tất lượt đếm gần nhất cho dòng hàng.",
    "discrepancy": "Giá trị chênh lệch giữa số lượng thực tế và số lượng hệ thống kỳ vọng.",
    "condition": "Tình trạng hàng hóa tại thời điểm nhập, xuất hoặc kiểm kê.",
    "has_discrepancy": "Cho biết dòng kiểm kê có chênh lệch cần xem xét hay không.",
    "recount_count": "Số lần dòng hàng đã được yêu cầu hoặc thực hiện kiểm đếm lại.",
    "last_recount_at": "Thời điểm kiểm đếm lại gần nhất.",
    "discrepancy_reason": "Lý do được chọn để giải thích chênh lệch số lượng.",
    "discrepancy_note": "Ghi chú chi tiết bổ sung cho chênh lệch số lượng.",
    "movement_delta_before_count": "Tổng biến động tồn được ghi nhận trước thời điểm kiểm đếm và cần tính vào số kỳ vọng.",
    "movement_delta_after_count": "Tổng biến động tồn xảy ra sau thời điểm kiểm đếm, dùng để đối chiếu với số hiện tại.",
    "evidence_urls": "Danh sách đường dẫn ảnh hoặc tài liệu làm bằng chứng cho kết quả xử lý.",
    "base_atp": "Số lượng tồn khả dụng nền dùng làm mốc tính toán chênh lệch kiểm kê.",
    "movement_detected": "Cho biết hệ thống có phát hiện biến động tồn trong khoảng thời gian kiểm kê hay không.",
}

EXACT_DESCRIPTIONS.update({
    "accent_color": "Mã màu chủ đạo dùng khi hiển thị chiến dịch voucher.",
    "accounting_category": "Nhóm hạch toán dùng để phân loại giao dịch điều chỉnh thành viên.",
    "account_name": "Tên chủ tài khoản nhận tiền chuyển khoản.",
    "account_number": "Số tài khoản ngân hàng nhận tiền chuyển khoản.",
    "accrual_day_of_month": "Ngày trong tháng mà hệ thống cộng phép định kỳ cho nhân viên.",
    "acted_at": "Thời điểm người được giao thực hiện quyết định duyệt.",
    "acted_by": "Mã người dùng đã thực hiện quyết định duyệt.",
    "action": "Loại thao tác được ghi vào nhật ký, ví dụ tạo mới, cập nhật, phê duyệt hoặc xóa.",
    "action_url": "Đường dẫn mở màn hình hoặc đối tượng liên quan khi người dùng chọn thông báo.",
    "active_version_id": "Mã phiên bản quyền truy cập đang được áp dụng cho người dùng.",
    "actual_quantity": "Số lượng thực tế nhận, đếm hoặc ghi nhận tại thời điểm thực hiện nghiệp vụ.",
    "address": "Địa chỉ của tổ chức hoặc cơ sở.",
    "after_sales_text": "Nội dung chính sách hoặc hướng dẫn sau bán hàng được in trên hóa đơn POS.",
    "amount": "Số tiền điều chỉnh hoặc bồi hoàn cho thành viên.",
    "amount_before_tax": "Tổng giá trị hàng hóa trước thuế của đơn hàng.",
    "annual_cap_units": "Số đơn vị phép tối đa nhân viên được tích lũy trong một năm.",
    "app_version": "Phiên bản ứng dụng POS đang chạy trên thiết bị.",
    "applicable_standard": "Tiêu chuẩn kỹ thuật hoặc chất lượng áp dụng cho sản phẩm.",
    "approval_attempt": "Số thứ tự vòng phê duyệt của đơn; tăng khi quy trình duyệt được khởi tạo lại.",
    "approval_chain": "Danh sách các bước hoặc vai trò phải phê duyệt theo đúng thứ tự.",
    "approver_id": "Mã người dùng chịu trách nhiệm phê duyệt chứng từ.",
    "assigned_by": "Mã người dùng đã gán vai trò hoặc quyền này.",
    "assigned_to": "Mã người dùng được giao thực hiện nhiệm vụ quy trình.",
    "assignment": "Cấu hình xác định người dùng hoặc vai trò được giao xử lý cấp duyệt.",
    "atp_deducted": "Cho biết số lượng trên phiếu xuất đã được trừ khỏi tồn khả dụng (ATP) hay chưa.",
    "attachment_urls": "Danh sách đường dẫn đến tệp chứng từ hoặc bằng chứng đính kèm.",
    "auto_approve": "Cho biết đối tượng có được tự động phê duyệt khi thỏa điều kiện cấu hình hay không.",
    "auto_print_after_payment": "Cho biết POS có tự động in phiếu ngay sau khi thanh toán thành công hay không.",
    "available_units": "Số đơn vị phép còn có thể dùng để tạo đơn nghỉ.",
    "balance_allocations": "Chi tiết số đơn vị phép được trừ từ từng nhóm số dư khi xử lý đơn nghỉ.",
    "bank_bin": "Mã BIN/NAPAS của ngân hàng nhận tiền, dùng để tạo mã QR chuyển khoản.",
    "barcode": "Mã vạch dùng để quét và nhận biết sản phẩm.",
    "base_url": "Địa chỉ gốc của API meInvoice theo môi trường kết nối.",
    "board_position": "Thứ tự hiển thị vai trò trên bảng công việc hoặc danh sách quản trị.",
    "body": "Nội dung chi tiết của thông báo hiển thị cho người nhận.",
    "body_font_size_pt": "Cỡ chữ phần nội dung phiếu POS, tính bằng point.",
    "business_dates": "Danh sách ngày kinh doanh có đơn hàng được đưa vào tác vụ phát hành hóa đơn.",
    "buyer": "Thông tin người mua được ghi trên hóa đơn, như tên, địa chỉ và mã số thuế.",
    "calculation": "Kết quả tính tiền, thuế và các khoản điều chỉnh của đơn hàng hoặc hóa đơn.",
    "calculation_version": "Phiên bản thuật toán tính tiền/thuế đã tạo ra kết quả hiện tại.",
    "campaign_name": "Tên chiến dịch được sao chép vào mã voucher để hiển thị và giữ lịch sử.",
    "cancel_money": "Số tiền của phần hàng đã hủy trong đơn hàng.",
    "cancel_qty": "Số lượng hàng đã hủy của dòng bán hàng.",
    "carryover_expiry_day": "Ngày trong tháng mà phần phép chuyển từ năm trước hết hạn.",
    "carryover_expiry_month": "Tháng trong năm mà phần phép chuyển từ năm trước hết hạn.",
    "category_description": "Mô tả phạm vi và đặc điểm của danh mục sản phẩm.",
    "category_vat_mapping": "Bảng ánh xạ danh mục sản phẩm với mức thuế VAT dùng khi lập hóa đơn.",
    "category_name": "Tên danh mục sản phẩm được lưu kèm dòng báo cáo để hiển thị và giữ lịch sử.",
    "channel": "Kênh gửi thông báo, ví dụ hiển thị trong ứng dụng hoặc gửi thông báo đẩy.",
    "code_counts": "Thống kê số mã voucher theo từng trạng thái trong chiến dịch.",
    "code_hash": "Mã băm của mã đăng ký thiết bị; hệ thống không lưu mã dùng một lần ở dạng rõ.",
    "code_length": "Số ký tự của phần ngẫu nhiên trong mã voucher được sinh.",
    "color": "Màu dùng để nhận biết vai trò trên giao diện.",
    "completed_by": "Mã người dùng đã hoàn thành nhiệm vụ quy trình.",
    "config_fingerprint": "Dấu vân tay của cấu hình phát hành hóa đơn tại thời điểm tạo lần chạy.",
    "config_snapshot": "Bản chụp cấu hình nghiệp vụ được dùng khi tạo lệnh điều chuyển, giúp kết quả không thay đổi nếu cấu hình sau đó được sửa.",
    "coordinate": "Tọa độ địa lý của cơ sở dùng cho bản đồ và kiểm tra vị trí.",
    "counts": "Số liệu tổng hợp số mục theo trạng thái trong tác vụ phát hành hóa đơn.",
    "created_by_name": "Tên người tạo giao dịch tại thời điểm ghi nhận, dùng để hiển thị và đối soát.",
    "create_time": "Thời điểm đơn hàng được tạo trên hệ thống bán hàng nguồn.",
    "creator_id": "Mã người dùng lập chứng từ.",
    "credential_hash": "Mã băm thông tin xác thực của thiết bị POS, dùng để kiểm tra thiết bị mà không lưu bí mật dạng rõ.",
    "current_node_ids": "Danh sách các nút quy trình đang hoạt động và chờ xử lý.",
    "cursor": "Vị trí xử lý hiện tại của tác vụ nền, dùng để tiếp tục công việc sau khi tạm dừng hoặc thử lại.",
    "customer_invoice_request_id": "Mã yêu cầu khách hàng cung cấp thông tin để nhận hóa đơn.",
    "customer_invoice_request_status": "Trạng thái yêu cầu xuất hóa đơn do khách hàng gửi.",
    "customer_invoice_request_submitted_at": "Thời điểm khách hàng gửi yêu cầu xuất hóa đơn.",
    "customer_name": "Tên khách hàng được ghi nhận trên đơn hàng.",
    "customer_phone": "Số điện thoại khách hàng tại thời điểm tạo đơn POS.",
    "days": "Danh sách ngày hoặc buổi nghỉ được đề nghị trong đơn nghỉ phép.",
    "deactivated_at": "Thời điểm tài khoản bị vô hiệu hóa.",
    "deactivated_by": "Mã người dùng đã vô hiệu hóa tài khoản.",
    "decision_reason": "Lý do người duyệt chấp thuận hoặc từ chối nhiệm vụ duyệt.",
    "default_buyer_address": "Địa chỉ người mua mặc định dùng khi đơn hàng không cung cấp thông tin này.",
    "default_buyer_name": "Tên người mua mặc định dùng khi đơn hàng không cung cấp thông tin này.",
    "default_buyer_tax_code": "Mã số thuế người mua mặc định dùng khi đơn hàng không cung cấp thông tin này.",
    "default_payment_method_name": "Tên phương thức thanh toán mặc định ghi trên hóa đơn khi không ánh xạ được dữ liệu nguồn.",
    "default_tax_rate": "Thuế suất mặc định áp dụng trên POS khi sản phẩm không có cấu hình riêng.",
    "default_unit_name": "Tên đơn vị tính mặc định ghi trên hóa đơn khi sản phẩm không có ánh xạ riêng.",
    "default_vat_rate_name": "Tên mức thuế VAT mặc định theo danh mục của meInvoice.",
    "delta": "Số đơn vị phép tăng hoặc giảm trong bút toán; giá trị dương là cộng phép, giá trị âm là trừ phép.",
    "dimensions": "Kích thước vật lý của sản phẩm theo cấu trúc đơn vị được hệ thống quy định.",
    "disabled_at": "Thời điểm mã nhận thông báo đẩy bị vô hiệu hóa.",
    "disabled_group_keys": "Danh sách khóa nhóm sản phẩm bị ẩn trên POS.",
    "disabled_product_ids": "Danh sách mã sản phẩm bị ẩn trên POS.",
    "discount_money": "Tổng số tiền giảm giá của đơn hàng hoặc dòng báo cáo.",
    "dispatched_at": "Thời điểm hàng được xác nhận rời cơ sở xuất.",
    "display_name": "Tên thân thiện dùng để nhận biết cấu hình hoặc tài khoản trên giao diện quản trị.",
    "display_order": "Thứ tự hiển thị sản phẩm trong ô/vị trí hoặc danh sách.",
    "distributed_at": "Thời điểm mã voucher được phân phối cho người nhận.",
    "distributed_to_phone": "Số điện thoại đã nhận mã voucher.",
    "due_at": "Thời hạn phải hoàn thành nhiệm vụ quy trình.",
    "duration_seconds": "Thời lượng phát của tệp video, tính bằng giây.",
    "edges": "Danh sách liên kết và điều kiện chuyển giữa các nút của phiên bản quy trình.",
    "edited_at": "Thời điểm chứng từ hóa đơn được chỉnh sửa thủ công gần nhất.",
    "edited_by": "Mã người dùng chỉnh sửa chứng từ hóa đơn gần nhất.",
    "eligible_source_order_ids": "Danh sách mã đơn hàng nguồn đủ điều kiện để phát hành hóa đơn.",
    "email_otp": "Mã OTP gửi qua email đang chờ xác minh; phải được bảo vệ và chỉ có hiệu lực ngắn hạn.",
    "email_otp_expires_at": "Thời điểm mã OTP email hết hiệu lực.",
    "emailed_at": "Thời điểm voucher được gửi qua email.",
    "emailed_to": "Địa chỉ email đã nhận voucher.",
    "enrolled_at": "Thời điểm thiết bị POS hoàn tất đăng ký với cơ sở.",
    "enrolled_by": "Mã người dùng đã phê duyệt hoặc thực hiện đăng ký thiết bị POS.",
    "entity_name": "Tên đối tượng nghiệp vụ tại thời điểm ghi nhật ký, giúp người đọc nhận biết bản ghi bị tác động.",
    "entity_type": "Loại đối tượng nghiệp vụ mà cấu hình, nhật ký hoặc quy trình áp dụng.",
    "entry_type": "Loại bút toán phép, ví dụ cộng định kỳ, sử dụng, hoàn trả, điều chỉnh hoặc hết hạn.",
    "environment": "Môi trường meInvoice được kết nối, ví dụ thử nghiệm hoặc sản xuất.",
    "excluded": "Danh sách đơn hàng bị loại khỏi lần phát hành kèm lý do không đủ điều kiện.",
    "expected_quantity": "Số lượng theo kế hoạch, chứng từ hoặc hệ thống trước khi ghi nhận số thực tế.",
    "expired_units": "Số đơn vị phép đã hết hạn và không còn sử dụng được.",
    "export_type": "Mục đích/loại xuất hàng, dùng để áp dụng quy trình và chứng từ phù hợp.",
    "facility_type": "Loại cơ sở được cấp quyền, dùng để diễn giải phạm vi và áp dụng đúng tập quyền.",
    "financially_edited": "Cho biết các giá trị tài chính của chứng từ hóa đơn đã được người dùng chỉnh sửa so với dữ liệu nguồn hay chưa.",
    "fingerprint_hash": "Mã băm dấu vân tay thiết bị, dùng để nhận biết thiết bị POS mà không lưu trực tiếp đặc điểm nhận dạng.",
    "fixed_transfer_only": "Cho biết thiết bị chỉ được tạo mã chuyển khoản đúng bằng số tiền đơn hàng hay có thể cho phép nhập số tiền khác.",
    "font_weight": "Độ đậm chữ dùng khi in phiếu POS.",
    "font_weights": "Cấu hình độ đậm chữ cho từng phần của hóa đơn POS.",
    "footer_message": "Thông điệp in ở cuối hóa đơn hoặc phiếu POS.",
    "generation_mode": "Cách tác vụ sinh hoặc mở rộng tập mã voucher.",
    "go_live_at": "Thời điểm cấu hình hóa đơn của cơ sở bắt đầu được sử dụng trong môi trường vận hành.",
    "goods_name": "Tên sản phẩm/hàng hóa được lưu kèm dòng báo cáo bán hàng.",
    "goods_type_name": "Tên loại hàng hóa được lưu kèm dòng báo cáo bán hàng.",
    "has_client_id": "Cho biết tài khoản meInvoice đã được cấu hình Client ID hay chưa; không lưu giá trị bí mật trong trường này.",
    "has_client_secret": "Cho biết tài khoản meInvoice đã được cấu hình Client Secret hay chưa; không lưu giá trị bí mật trong trường này.",
    "has_password": "Cho biết tài khoản meInvoice đã được cấu hình mật khẩu hay chưa; không lưu mật khẩu trong trường này.",
    "has_username": "Cho biết tài khoản meInvoice đã được cấu hình tên đăng nhập hay chưa.",
    "height": "Chiều cao theo pixel của ảnh hoặc video.",
    "held_units": "Số đơn vị phép đang được giữ cho các đơn chờ duyệt và chưa thể dùng cho đơn khác.",
    "hk_order_number": "Số đơn hàng trên hệ thống HK dùng để đối soát với đơn và hóa đơn nội bộ.",
    "holiday_date": "Ngày nghỉ được công ty công bố.",
    "hotline": "Số điện thoại hỗ trợ được in trên hóa đơn POS.",
    "hs_code": "Mã HS dùng để phân loại hàng hóa trong nghiệp vụ thuế và hải quan.",
    "image_storage_path": "Đường dẫn nội bộ của ảnh chiến dịch trong dịch vụ lưu trữ.",
    "image_url": "Đường dẫn hiển thị ảnh chiến dịch hoặc hình ảnh của đối tượng.",
    "inserted_count": "Số đơn hàng mới được thêm trong lần đồng bộ.",
    "instructions": "Nội dung hướng dẫn khách hàng được in trên phiếu POS.",
    "inv_series": "Ký hiệu mẫu/sê-ri hóa đơn được dùng khi phát hành trên meInvoice.",
    "invoice_code": "Mã tra cứu hoặc mã do meInvoice trả về cho hóa đơn đã phát hành.",
    "invoice_date_source": "Quy tắc chọn ngày hóa đơn từ ngày kinh doanh, thời điểm thanh toán hoặc nguồn được cấu hình.",
    "invoice_document_source_payload_hash": "Mã băm dữ liệu đơn hàng nguồn đã dùng để tạo chứng từ hóa đơn gần nhất.",
    "invoice_document_stale": "Cho biết chứng từ hóa đơn hiện tại đã cũ do dữ liệu đơn hàng nguồn thay đổi hay chưa.",
    "invoice_document_status": "Trạng thái chứng từ hóa đơn liên kết với đơn hàng nguồn.",
    "invoice_kind": "Loại chứng từ hóa đơn, xác định nội dung và quy trình phát hành.",
    "invoice_number": "Số hóa đơn chính thức do meInvoice cấp sau khi phát hành thành công.",
    "invoice_qr_hint_font_size_pt": "Cỡ chữ phần hướng dẫn dưới mã QR yêu cầu hóa đơn, tính bằng point.",
    "invoice_qr_size_mm": "Kích thước mã QR yêu cầu hóa đơn trên bản in, tính bằng milimét.",
    "invoice_qr_title_font_size_pt": "Cỡ chữ tiêu đề mã QR yêu cầu hóa đơn, tính bằng point.",
    "invoice_with_code": "Cho biết cơ sở phát hành hóa đơn điện tử có mã của cơ quan thuế hay không.",
    "ip_address": "Địa chỉ IP của thiết bị tại thời điểm thực hiện chấm công hoặc thao tác được ghi nhật ký.",
    "is_invoice_calculating_machine": "Cho biết hóa đơn được phát hành từ máy tính tiền theo cấu hình của cơ quan thuế hay không.",
    "is_read": "Cho biết người nhận đã mở hoặc đánh dấu thông báo là đã đọc hay chưa.",
    "is_serialized": "Cho biết sản phẩm có bắt buộc quản lý từng đơn vị theo số sê-ri hay không.",
    "is_valid": "Cho biết dòng dữ liệu nhập đã vượt qua toàn bộ kiểm tra bắt buộc hay chưa.",
    "issue_deadline_at": "Hạn cuối phải phát hành hóa đơn cho đơn hàng hoặc tác vụ.",
    "issue_eligible": "Cho biết chứng từ đã đáp ứng các điều kiện bắt buộc để được phát hành hay chưa.",
    "issue_retry_eligible": "Cho biết lỗi phát hành gần nhất có cho phép hệ thống thử phát hành lại hay không.",
    "issue_scope": "Phạm vi đơn hàng mà cấu hình cho phép phát hành hóa đơn.",
    "issue_type": "Loại sai lệch hoặc vấn đề chất lượng được báo cáo.",
    "item_count": "Số dòng hàng trong đơn hoặc kết quả báo cáo.",
    "item_name_mapping": "Quy tắc ánh xạ tên hàng từ dữ liệu bán hàng sang nội dung hóa đơn.",
    "item_unit_mapping": "Quy tắc ánh xạ đơn vị tính của từng mặt hàng sang danh mục meInvoice.",
    "items": "Danh sách dòng hàng và thông tin số lượng, đơn giá, thuế của chứng từ hoặc đơn hàng.",
    "label": "Tên hiển thị của cấp duyệt để người dùng nhận biết nhiệm vụ.",
    "last_error_message": "Thông báo chi tiết của lỗi gần nhất khi xử lý tác vụ voucher.",
    "last_issue_error": "Thông báo lỗi gần nhất khi phát hành chứng từ hóa đơn.",
    "last_issue_error_code": "Mã lỗi gần nhất khi phát hành hóa đơn, dùng để xác định nguyên nhân và khả năng thử lại.",
    "last_seen_at": "Thời điểm thiết bị hoặc mã nhận thông báo gần nhất kết nối với hệ thống.",
    "last_template_sync_at": "Thời điểm danh sách mẫu hóa đơn được đồng bộ từ meInvoice gần nhất.",
    "last_test_error_code": "Mã lỗi của lần kiểm tra kết nối meInvoice gần nhất.",
    "last_test_succeeded": "Kết quả thành công/thất bại của lần kiểm tra kết nối meInvoice gần nhất.",
    "last_tested_at": "Thời điểm tài khoản meInvoice được kiểm tra kết nối gần nhất.",
    "leave_date": "Ngày nghỉ được đặt chỗ theo buổi để ngăn tạo đơn trùng.",
    "leave_year": "Năm phép mà số dư hoặc bút toán được hạch toán.",
    "legacy_metadata": "Dữ liệu bổ sung được giữ lại từ phiên bản hoặc hệ thống cũ để tương thích và truy vết.",
    "legacy_status": "Trạng thái đơn theo mô hình dữ liệu POS cũ, được giữ để tương thích.",
    "level": "Số thứ tự cấp duyệt trong chuỗi phê duyệt.",
    "levels": "Danh sách các cấp duyệt và quy tắc xác định người chịu trách nhiệm ở từng cấp.",
    "logo_checksum_sha256": "Mã băm SHA-256 của tệp logo, dùng để kiểm tra toàn vẹn và nhận biết thay đổi.",
    "logo_content_type": "Loại MIME của tệp logo.",
    "logo_content_url": "Đường dẫn tải nội dung logo để in trên POS.",
    "logo_contrast_percent": "Mức điều chỉnh tương phản của logo khi in, tính theo phần trăm.",
    "logo_data_url": "Nội dung logo mã hóa dạng data URL dùng trực tiếp khi in; có thể để trống nếu dùng tệp lưu trữ.",
    "logo_file_size_bytes": "Dung lượng tệp logo, tính bằng byte.",
    "logo_max_height_mm": "Chiều cao tối đa của logo trên bản in, tính bằng milimét.",
    "logo_storage_path": "Đường dẫn nội bộ của tệp logo trong dịch vụ lưu trữ.",
    "logo_width_mm": "Chiều rộng logo trên bản in, tính bằng milimét.",
    "manager_id": "Mã người dùng được chỉ định là người quản lý cơ sở.",
    "manufacturer": "Tên nhà sản xuất sản phẩm.",
    "manufacturer_address": "Địa chỉ của nhà sản xuất sản phẩm.",
    "mapped_payment_method": "Phương thức thanh toán sau khi ánh xạ dữ liệu nguồn sang danh mục dùng cho hóa đơn.",
    "mapping": "Cấu hình ánh xạ vùng ô trong mẫu báo cáo với các trường dữ liệu hệ thống.",
    "mapping_version": "Phiên bản bộ quy tắc ánh xạ đã được dùng để chuẩn hóa dữ liệu.",
    "match_status": "Kết quả đối chiếu đơn hàng nguồn với dữ liệu nội bộ hoặc chứng từ hóa đơn.",
    "member_code": "Mã thành viên trên hệ thống bên ngoài.",
    "member_name": "Tên thành viên tại thời điểm ghi nhận giao dịch.",
    "member_uid": "Định danh kỹ thuật của thành viên trên hệ thống bên ngoài.",
    "member_warehouse_ids": "Danh sách mã cơ sở thành viên được gom vào cùng một ánh xạ cửa hàng bên ngoài.",
    "message": "Thông điệp chính được in trên phiếu rút thăm hoặc trả về từ nghiệp vụ.",
    "mfa_enabled": "Cho biết tài khoản đã bật xác thực đa yếu tố hay chưa.",
    "mfa_secret": "Bí mật dùng để tạo mã xác thực đa yếu tố; trường này phải được mã hóa và không hiển thị trực tiếp.",
    "misa_error_code": "Mã lỗi do MISA meInvoice trả về trong lần phát hành gần nhất.",
    "mode": "Chế độ ánh xạ cửa hàng ngoài với cơ sở, ví dụ một-một hoặc gom nhiều cơ sở.",
    "monthly_accrual_units": "Số đơn vị phép được cộng cho nhân viên trong mỗi kỳ hằng tháng.",
    "new_value": "Giá trị của đối tượng sau khi thao tác thay đổi được thực hiện.",
    "next_assignment": "Thông tin người dùng hoặc vai trò nhận nhiệm vụ sau khi chuyển giao.",
    "next_attempt_at": "Thời điểm sớm nhất hệ thống được phép thử xử lý lại dòng phát hành.",
    "node_type": "Loại nút quy trình, quyết định cách giao việc và xử lý khi nút được kích hoạt.",
    "nodes": "Danh sách các bước/nút và cấu hình của từng bước trong phiên bản quy trình.",
    "normalized_items": "Danh sách dòng hàng sau khi chuẩn hóa từ dữ liệu đơn hàng nguồn.",
    "normalized_phone": "Số điện thoại khách hàng đã chuẩn hóa để tìm kiếm và đối chiếu.",
    "old_value": "Giá trị của đối tượng trước khi thao tác thay đổi được thực hiện.",
    "operating_system": "Tên và phiên bản hệ điều hành của thiết bị POS.",
    "order_id": "Mã đơn hàng trên hệ thống nguồn dùng để liên kết và đối soát dữ liệu báo cáo.",
    "operator_id": "Mã người dùng hoặc nhân viên thực hiện đơn hàng POS.",
    "option_user_defined": "Các trường mở rộng do người dùng cấu hình để gửi sang meInvoice.",
    "order_count": "Tổng số đơn hàng được xử lý trong lần đồng bộ.",
    "order_number": "Số chứng từ hoặc đơn hàng hiển thị cho người dùng và dùng để tra cứu.",
    "organization_image_url": "Đường dẫn hình ảnh hoặc logo đại diện của tổ chức.",
    "original_money": "Tổng tiền nguyên gốc nhận từ hệ thống bán hàng trước khi hệ thống tính lại.",
    "output_checksum": "Mã kiểm tra của tệp kết quả, dùng để xác minh tệp không bị thay đổi.",
    "output_storage_path": "Đường dẫn lưu tệp kết quả do tác vụ voucher tạo ra.",
    "package_ticket_counts": "Cấu hình số phiếu rút thăm được in tương ứng với từng gói hoặc mức mua hàng.",
    "paid_at": "Thời điểm đơn POS được thanh toán thành công.",
    "paper_size": "Khổ giấy dùng để in hóa đơn hoặc phiếu POS.",
    "password_hash": "Mã băm mật khẩu dùng để xác thực đăng nhập; hệ thống không lưu mật khẩu dạng rõ.",
    "payment_method": "Phương thức thanh toán nhận từ đơn hàng nguồn.",
    "payment_method_mapping": "Bảng ánh xạ phương thức thanh toán nguồn sang tên phương thức ghi trên hóa đơn.",
    "payment_method_name": "Tên phương thức thanh toán sẽ được ghi trên chứng từ hóa đơn.",
    "payment_time": "Thời điểm đơn hàng được thanh toán trên hệ thống nguồn.",
    "payment_status": "Trạng thái thanh toán của đơn POS.",
    "pay_method": "Tên phương thức thanh toán trong dữ liệu báo cáo bán hàng.",
    "pending_probation_units": "Số đơn vị phép đã phát sinh nhưng bị khóa cho đến khi nhân viên hoàn tất thử việc.",
    "permission": "Trạng thái quyền nhận thông báo đẩy do trình duyệt hoặc thiết bị cấp.",
    "picked_quantity": "Số lượng thực tế đã được lấy khỏi vị trí để thực hiện phiếu xuất.",
    "platform": "Nền tảng của thiết bị nhận thông báo, ví dụ web, iOS hoặc Android.",
    "playlist": "Danh sách media và thứ tự phát trên màn hình khách hàng POS.",
    "pos_order_status": "Trạng thái đơn hàng theo hệ thống POS nguồn.",
    "posting_date": "Ngày hạch toán bút toán tăng/giảm số dư phép.",
    "prefix": "Chuỗi ký tự cố định đặt ở đầu mã voucher.",
    "preflight": "Kết quả kiểm tra trước khi tạo hóa đơn, gồm điều kiện hợp lệ và các vấn đề cần xử lý.",
    "prepared_payload_hash": "Mã băm của dữ liệu đã chuẩn bị gửi meInvoice, dùng để bảo đảm lần thử lại phát hành đúng cùng nội dung.",
    "preview_fingerprint": "Dấu vân tay của kết quả xem trước danh sách đơn được chọn, dùng để phát hiện dữ liệu thay đổi trước khi tạo tác vụ.",
    "previous_assignment": "Thông tin người dùng hoặc vai trò chịu trách nhiệm trước khi nhiệm vụ được chuyển giao.",
    "price": "Đơn giá bán của sản phẩm trong dòng báo cáo.",
    "price_includes_vat": "Cho biết giá bán nguồn đã bao gồm VAT hay chưa để hệ thống tính tiền thuế đúng.",
    "priority": "Mức ưu tiên của thông báo, dùng để sắp xếp hoặc chọn cách hiển thị.",
    "probation_accrual_locked": "Cho biết phép phát sinh trong thời gian thử việc có bị khóa sử dụng đến khi nhân viên chính thức hay không.",
    "product_font_size_pt": "Cỡ chữ tên sản phẩm trên phiếu POS, tính bằng point.",
    "product_image_url": "Đường dẫn hình ảnh đại diện của sản phẩm.",
    "product_material": "Vật liệu hoặc thành phần chính của sản phẩm.",
    "product_origin": "Quốc gia hoặc nơi xuất xứ của sản phẩm.",
    "product_type": "Loại sản phẩm, dùng để áp dụng cách quản lý và xử lý phù hợp.",
    "product_names": "Danh sách tên sản phẩm được lưu trong bản tóm tắt đơn để tìm kiếm nhanh.",
    "program_name": "Tên chương trình rút thăm được in trên phiếu.",
    "progress": "Thông tin tiến độ tác vụ, gồm số mục đã xử lý, thành công hoặc thất bại.",
    "published_at": "Thời điểm phiên bản quy trình được công bố để sử dụng.",
    "published_by": "Mã người dùng công bố phiên bản quy trình.",
    "purpose": "Mục đích nghiệp vụ của bản ghi hoặc lần xử lý.",
    "qr_size_mm": "Kích thước mã QR trên phiếu POS, tính bằng milimét.",
    "qty": "Số lượng sản phẩm bán ra trong dòng báo cáo.",
    "quantity_affected": "Số lượng sản phẩm bị ảnh hưởng bởi sai lệch hoặc vấn đề chất lượng.",
    "quarantine_reason": "Lý do sản phẩm phải được đưa vào trạng thái cách ly.",
    "quarantined_at": "Thời điểm sản phẩm được đưa vào cách ly.",
    "read_at": "Thời điểm người nhận đọc thông báo lần đầu.",
    "real_money": "Số tiền thực thu của đơn hàng sau giảm giá và điều chỉnh.",
    "reassigned_by": "Mã người dùng thực hiện chuyển giao nhiệm vụ duyệt.",
    "reauth_confirmed_at": "Thời điểm người dùng hoàn tất xác thực lại cho lệnh điều chuyển.",
    "reauth_confirmed_by": "Mã người dùng đã xác thực lại lệnh điều chuyển.",
    "received_at": "Thời điểm cơ sở nhận xác nhận đã nhận hàng điều chuyển.",
    "received_by": "Mã người dùng xác nhận nhận hàng tại cơ sở đích.",
    "received_quantity": "Số lượng thực tế cơ sở đích đã nhận của dòng điều chuyển.",
    "recipient_department": "Bộ phận nhận hàng trên phiếu xuất.",
    "recipient_email": "Địa chỉ email nhận tệp hoặc danh sách voucher.",
    "recipient_name": "Tên người hoặc đơn vị nhận hàng trên phiếu xuất.",
    "record_type": "Loại dữ liệu của dòng nhập, ví dụ số dư phép hoặc bút toán điều chỉnh.",
    "ref_id": "Mã tham chiếu duy nhất gửi sang meInvoice để đối soát và chống phát hành trùng.",
    "reference_type": "Loại đối tượng/chứng từ được phiếu xuất tham chiếu.",
    "rejected_at": "Thời điểm chứng từ hoặc yêu cầu bị từ chối.",
    "rejected_by": "Mã người dùng từ chối chứng từ hoặc yêu cầu.",
    "release_notes": "Ghi chú giải thích điều kiện và kết quả giải phóng hàng khỏi cách ly.",
    "released_at": "Thời điểm hàng được giải phóng khỏi trạng thái cách ly.",
    "released_by": "Mã người dùng phê duyệt giải phóng hàng khỏi cách ly.",
    "remote_code": "Mã kết quả do hệ thống thành viên bên ngoài trả về.",
    "remote_message": "Thông báo kết quả do hệ thống thành viên bên ngoài trả về.",
    "remote_total_value": "Tổng giá trị thành viên trên hệ thống bên ngoài sau khi giao dịch hoàn tất.",
    "report_number": "Số báo cáo không phù hợp hiển thị cho người dùng và dùng để tra cứu.",
    "request_type": "Loại nghỉ hoặc loại yêu cầu, dùng để áp dụng đúng quy tắc số dư và phê duyệt.",
    "require_evidence": "Cho biết quy trình có bắt buộc đính kèm bằng chứng trước khi hoàn tất hay không.",
    "require_otp": "Cho biết quy trình có bắt buộc xác thực OTP ở bước phê duyệt hay không.",
    "requires_evidence": "Cho biết báo cáo phải có ảnh hoặc tài liệu bằng chứng trước khi được xử lý hay không.",
    "requires_reauth": "Cho biết lệnh điều chuyển có bắt buộc người dùng xác thực lại trước thao tác nhạy cảm hay không.",
    "resolution_notes": "Ghi chú chi tiết về biện pháp đã thực hiện để xử lý sai lệch hoặc vấn đề chất lượng.",
    "resolution_type": "Loại biện pháp xử lý báo cáo không phù hợp.",
    "resolved_by": "Mã người dùng chịu trách nhiệm xử lý hoàn tất báo cáo không phù hợp.",
    "retry_count": "Số lần hệ thống đã thực hiện lại sau lần xử lý đầu tiên.",
    "retry_eligible": "Cho biết lỗi hiện tại có thuộc nhóm được phép tự động thử lại hay không.",
    "review_note": "Ghi chú của người xem xét chứng từ hóa đơn.",
    "revoke_reason": "Lý do mã voucher hoặc thiết bị bị thu hồi.",
    "revoked_at": "Thời điểm mã, thiết bị hoặc voucher bị thu hồi.",
    "revoked_by": "Mã người dùng thực hiện thu hồi.",
    "reward_type": "Loại ưu đãi của voucher, ví dụ giảm theo số tiền hoặc tỷ lệ phần trăm.",
    "reward_value": "Giá trị ưu đãi tương ứng với loại voucher.",
    "role_name": "Tên vai trò được sao chép tại thời điểm gán để hiển thị và giữ lịch sử.",
    "scope_origin": "Nguồn hình thành phạm vi vai trò, ví dụ gán trực tiếp, nơi làm việc hoặc kế thừa chính sách.",
    "scope_warehouse_ids": "Danh sách mã cơ sở mà định nghĩa quy trình được phép áp dụng.",
    "selection_mode": "Cách người dùng chọn đơn hàng cho lần phát hành hàng loạt, ví dụ chọn tất cả đủ điều kiện hoặc chọn thủ công.",
    "seller_shop_code": "Mã cửa hàng của bên bán gửi sang meInvoice.",
    "seller_shop_name": "Tên cửa hàng của bên bán ghi trên dữ liệu hóa đơn.",
    "session_token": "Mã phiên đăng nhập tại thời điểm thao tác, dùng để điều tra và truy vết bảo mật.",
    "show_cashier": "Cho biết có in tên thu ngân trên hóa đơn POS hay không.",
    "show_contact": "Cho biết có in địa chỉ và thông tin liên hệ của cơ sở hay không.",
    "show_invoice_request_qr": "Cho biết có in mã QR để khách hàng gửi yêu cầu xuất hóa đơn hay không.",
    "show_issued_at": "Cho biết có in thời điểm phát hành phiếu POS hay không.",
    "show_item_tax": "Cho biết có hiển thị thuế của từng dòng hàng trên hóa đơn POS hay không.",
    "show_logo": "Cho biết có in logo trên hóa đơn hoặc phiếu POS hay không.",
    "show_order_code": "Cho biết có in mã đơn hàng trên phiếu POS hay không.",
    "show_price": "Cho biết có in giá trị tiền trên phiếu POS hay không.",
    "show_sequence": "Cho biết có in số thứ tự phục vụ trên phiếu POS hay không.",
    "show_theme_message": "Cho biết có in thông điệp theo chủ đề trên hóa đơn POS hay không.",
    "sign_type": "Hình thức ký số được meInvoice sử dụng khi phát hành hóa đơn.",
    "sku_mapping": "Bảng ánh xạ mã sản phẩm nội bộ với mã hàng gửi sang meInvoice.",
    "sort_order": "Thứ tự sắp xếp ô/vị trí trong khu vực.",
    "source_account_key": "Khóa tài khoản hoặc cửa hàng trên hệ thống nguồn dùng để nhận biết ánh xạ.",
    "source_action_time": "Thời điểm nghiệp vụ xảy ra trên hệ thống nguồn.",
    "source_create_time": "Thời điểm đơn hàng được tạo trên hệ thống nguồn.",
    "source_entity_type": "Loại đối tượng nguồn đã phát sinh thông báo.",
    "source_financial_fingerprint": "Dấu vân tay của các giá trị tài chính từ đơn hàng nguồn, dùng để phát hiện thay đổi sau khi tạo chứng từ.",
    "source_hash": "Mã băm của dữ liệu nguồn dùng để phát hiện voucher hoặc nội dung bị thay đổi.",
    "source_id": "Mã chứng từ hoặc đối tượng nguồn làm phát sinh báo cáo không phù hợp.",
    "source_order_number": "Số đơn hàng nguồn được lưu trên chứng từ hóa đơn để đối soát.",
    "source_payload_hash": "Mã băm của toàn bộ dữ liệu đơn hàng nguồn, dùng để phát hiện thay đổi và tránh xử lý lặp.",
    "source_status": "Trạng thái đơn hàng nhận từ hệ thống nguồn.",
    "source_sync_time": "Thời điểm dữ liệu được đồng bộ tại hệ thống nguồn.",
    "source_system": "Tên hệ thống bên ngoài cung cấp đơn hàng hoặc dữ liệu được đồng bộ.",
    "source_type": "Loại nghiệp vụ nguồn làm phát sinh báo cáo, ví dụ nhập, xuất hoặc kiểm kê.",
    "started_by": "Mã người dùng khởi tạo lần thực thi quy trình.",
    "status_effective_date": "Ngày trạng thái tài khoản bắt đầu có hiệu lực.",
    "status_reason": "Lý do tài khoản chuyển sang trạng thái hiện tại.",
    "status_source_id": "Mã đối tượng nghiệp vụ làm phát sinh thay đổi trạng thái tài khoản.",
    "status_label": "Tên trạng thái đã được chuyển sang ngôn ngữ hiển thị trong báo cáo.",
    "step_options": "Các tùy chọn chi tiết áp dụng cho từng bước của quy trình phê duyệt.",
    "store_address": "Địa chỉ cơ sở được in trên hóa đơn hoặc phiếu POS.",
    "store_name": "Tên cơ sở được in trên hóa đơn hoặc phiếu POS.",
    "stored_category": "Nhóm giao dịch do hệ thống thành viên bên ngoài lưu nhận.",
    "subtitle": "Dòng phụ đề được in dưới tiêu đề phiếu POS.",
    "suffix": "Chuỗi ký tự cố định đặt ở cuối mã voucher.",
    "summary": "Số liệu tổng hợp về tập đơn được chọn, đủ điều kiện, bị loại và đã tạo tác vụ.",
    "supplier_name": "Tên nhà cung cấp giao hàng theo phiếu nhập.",
    "sync_status": "Trạng thái đồng bộ của đơn POS với hệ thống phía máy chủ hoặc hệ thống ngoài.",
    "system_money": "Tổng tiền do hệ thống hiện tại tính lại từ các dòng hàng.",
    "sys_money": "Tổng tiền hàng theo dữ liệu hệ thống bán hàng nguồn.",
    "target_valid_to": "Ngày hết hiệu lực mới mà tác vụ gia hạn voucher cần áp dụng.",
    "tax_code": "Mã số thuế của tổ chức, pháp nhân hoặc người mua.",
    "tax_money": "Tổng số tiền thuế của đơn hàng hoặc dòng báo cáo.",
    "tax_rate_source": "Nguồn xác định thuế suất, ví dụ theo sản phẩm, danh mục hoặc giá trị mặc định.",
    "technical_specifications": "Các thông số kỹ thuật dùng để mô tả và kiểm tra sản phẩm.",
    "template_key": "Khóa mẫu nội dung dùng để tạo tiêu đề và nội dung thông báo.",
    "template_params": "Các tham số được điền vào mẫu khi tạo nội dung thông báo.",
    "terminal_name": "Tên thiết bị đầu cuối/POS đã ghi nhận đơn hàng.",
    "theme": "Chủ đề trình bày của hóa đơn POS.",
    "theme_message_font_size_pt": "Cỡ chữ của thông điệp theo chủ đề, tính bằng point.",
    "theme_messages": "Danh sách thông điệp có thể in theo chủ đề trên hóa đơn POS.",
    "ticket_height_mm": "Chiều cao phiếu POS, tính bằng milimét.",
    "ticket_title": "Tiêu đề chính được in trên phiếu POS hoặc phiếu rút thăm.",
    "title_font_size_pt": "Cỡ chữ tiêu đề phiếu POS, tính bằng point.",
    "token": "Mã đăng ký nhận thông báo đẩy do trình duyệt hoặc nền tảng thiết bị cấp.",
    "total_issued": "Tổng số mã voucher đã được phát hành trong chiến dịch.",
    "total_price": "Thành tiền của dòng giao dịch, thường bằng số lượng nhân với đơn giá sau quy tắc làm tròn.",
    "total_units": "Tổng số đơn vị phép mà đơn nghỉ yêu cầu sử dụng.",
    "total_amount": "Tổng số tiền phải thanh toán của đơn POS.",
    "total_qty": "Tổng số lượng sản phẩm trong đơn hàng báo cáo.",
    "transfer_type": "Loại điều chuyển hàng, dùng để áp dụng quy trình nhận/xuất phù hợp.",
    "type": "Loại của {entity}, dùng để áp dụng quy tắc xử lý phù hợp.",
    "unchanged_count": "Số đơn hàng không thay đổi so với dữ liệu đã có trong lần đồng bộ.",
    "unit": "Đơn vị tính cơ bản của sản phẩm.",
    "unit_name_mapping": "Bảng ánh xạ tên đơn vị nội bộ sang tên đơn vị được meInvoice chấp nhận.",
    "updated_count": "Số đơn hàng hiện có đã được cập nhật trong lần đồng bộ.",
    "updated_by_uid": "Mã người dùng cập nhật cấu hình POS gần nhất.",
    "used_at": "Thời điểm mã voucher hoặc mã đăng ký được sử dụng.",
    "used_by_staff_id": "Mã nhân viên xác nhận sử dụng voucher.",
    "used_by_staff_name": "Tên nhân viên xác nhận sử dụng voucher tại thời điểm giao dịch.",
    "used_units": "Số đơn vị phép đã được sử dụng trong năm.",
    "user_agent": "Chuỗi nhận dạng trình duyệt/thiết bị đã đăng ký nhận thông báo.",
    "user_name": "Tên người dùng tại thời điểm thao tác, được lưu để hiển thị nhật ký.",
    "username": "Tên đăng nhập duy nhất của tài khoản.",
    "valid_to": "Thời điểm mã voucher hết hiệu lực.",
    "validated_at": "Thời điểm cấu hình được kiểm tra và xác nhận hợp lệ gần nhất.",
    "validation_error_code": "Mã lỗi của lần kiểm tra cấu hình gần nhất.",
    "validation_issues": "Danh sách lỗi hoặc cảnh báo phải xử lý trước khi chứng từ được phép phát hành.",
    "voucher_code_ids": "Danh sách mã định danh voucher được xử lý trong dòng tác vụ.",
    "voucher_number": "Số phiếu nhập hoặc phiếu xuất hiển thị cho người dùng và dùng để tra cứu.",
    "warehouse_description": "Mô tả chức năng, phạm vi hoạt động hoặc ghi chú của cơ sở.",
    "warehouse_image_url": "Đường dẫn hình ảnh đại diện của cơ sở.",
    "warehouse_location_description": "Mô tả chức năng hoặc đặc điểm của khu vực trong cơ sở.",
    "warehouse_location_image_url": "Đường dẫn hình ảnh đại diện của khu vực trong cơ sở.",
    "warehouse_name": "Tên cơ sở được lưu kèm dữ liệu báo cáo để hiển thị và giữ lịch sử.",
    "width": "Chiều rộng theo pixel của ảnh hoặc video.",
    "workbook_meta": "Thông tin cấu trúc tệp bảng tính mẫu, như tên trang tính, vùng dữ liệu và định dạng cần giữ.",
})


TOKEN_TRANSLATIONS = {
    "account": "tài khoản", "action": "thao tác", "active": "hoạt động", "actual": "thực tế",
    "address": "địa chỉ", "after": "sau", "allocation": "phân bổ", "allocations": "các phân bổ",
    "amount": "số tiền", "annual": "hằng năm", "app": "ứng dụng", "approval": "phê duyệt",
    "assigned": "được giao", "attachment": "tệp đính kèm", "attachments": "các tệp đính kèm",
    "auto": "tự động", "available": "khả dụng", "balance": "số dư", "bank": "ngân hàng",
    "before": "trước", "bin": "mã BIN", "body": "nội dung", "buyer": "người mua",
    "calculation": "kết quả tính toán", "campaign": "chiến dịch", "cancel": "hủy", "cancelled": "đã hủy",
    "category": "danh mục", "chain": "chuỗi", "channel": "kênh", "checksum": "mã kiểm tra",
    "client": "ứng dụng khách", "code": "mã", "color": "màu", "completed": "hoàn tất",
    "config": "cấu hình", "contact": "liên hệ", "content": "nội dung", "coordinate": "tọa độ",
    "count": "số lượng", "counts": "thống kê số lượng", "create": "tạo", "created": "được tạo",
    "current": "hiện hành", "customer": "khách hàng", "date": "ngày", "deadline": "hạn xử lý",
    "default": "mặc định", "delta": "mức thay đổi", "destination": "đích", "device": "thiết bị",
    "disabled": "bị ẩn/tắt", "discount": "giảm giá", "display": "hiển thị", "document": "tài liệu",
    "duration": "thời lượng", "edge": "quan hệ", "eligible": "đủ điều kiện", "email": "email",
    "employee": "nhân viên", "entity": "đối tượng", "environment": "môi trường", "evidence": "bằng chứng",
    "excluded": "bị loại", "expiry": "hết hạn", "external": "bên ngoài", "file": "tệp",
    "financial": "tài chính", "fingerprint": "dấu vân tay", "fixed": "cố định", "font": "phông chữ",
    "footer": "chân phiếu", "generation": "sinh dữ liệu", "goods": "hàng hóa", "group": "nhóm",
    "hash": "mã băm", "height": "chiều cao", "held": "đang giữ", "hint": "hướng dẫn",
    "history": "lịch sử", "holiday": "ngày nghỉ", "hotline": "số hotline", "image": "hình ảnh",
    "import": "nhập dữ liệu", "includes": "bao gồm", "instructions": "hướng dẫn", "invoice": "hóa đơn",
    "issued": "phát hành", "item": "dòng hàng", "items": "danh sách dòng hàng", "label": "nhãn hiển thị",
    "last": "gần nhất", "legacy": "hệ thống cũ", "level": "cấp", "logo": "logo",
    "mapping": "ánh xạ", "max": "tối đa", "member": "thành viên", "message": "thông điệp",
    "method": "phương thức", "min": "tối thiểu", "mode": "chế độ", "money": "số tiền",
    "monthly": "hằng tháng", "name": "tên", "next": "tiếp theo", "number": "số",
    "option": "tùy chọn", "order": "đơn hàng", "origin": "xuất xứ", "original": "gốc",
    "otp": "OTP", "paper": "khổ giấy", "password": "mật khẩu", "payment": "thanh toán",
    "percent": "tỷ lệ phần trăm", "permission": "quyền", "phone": "điện thoại", "platform": "nền tảng",
    "playlist": "danh sách phát", "position": "vị trí", "prefix": "tiền tố", "previous": "trước đó",
    "price": "giá", "priority": "mức ưu tiên", "product": "sản phẩm", "progress": "tiến độ",
    "purpose": "mục đích", "quantity": "số lượng", "qr": "mã QR", "rate": "thuế suất",
    "recipient": "người nhận", "reference": "tham chiếu", "remote": "hệ thống ngoài", "report": "báo cáo",
    "request": "yêu cầu", "required": "bắt buộc", "result": "kết quả", "retry": "thử lại",
    "reward": "ưu đãi", "role": "vai trò", "scope": "phạm vi", "secret": "bí mật xác thực",
    "seller": "người bán", "sequence": "thứ tự", "session": "phiên", "sha256": "SHA-256",
    "shop": "cửa hàng", "show": "hiển thị", "size": "kích thước", "source": "nguồn",
    "staff": "nhân viên", "status": "trạng thái", "storage": "lưu trữ", "store": "cơ sở bán hàng",
    "subtitle": "phụ đề", "suffix": "hậu tố", "summary": "tổng hợp", "supplier": "nhà cung cấp",
    "sync": "đồng bộ", "system": "hệ thống", "target": "mục tiêu", "tax": "thuế",
    "template": "mẫu", "terminal": "thiết bị đầu cuối", "ticket": "phiếu", "time": "thời gian",
    "token": "mã thiết bị", "total": "tổng", "transaction": "giao dịch", "transfer": "chuyển khoản/điều chuyển",
    "type": "loại", "unit": "đơn vị", "url": "đường dẫn", "urls": "các đường dẫn",
    "user": "người dùng", "used": "đã sử dụng", "username": "tên đăng nhập", "valid": "hợp lệ",
    "validation": "kiểm tra hợp lệ", "value": "giá trị", "version": "phiên bản", "voucher": "voucher",
    "warehouse": "cơ sở", "weight": "độ đậm", "width": "chiều rộng", "workbook": "tệp bảng tính",
}


def split_name(name: str) -> list[str]:
    snake = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", name).lower()
    return [part for part in snake.split("_") if part]


def canonical_name(name: str) -> str:
    return "_".join(split_name(name))


def humanize_field(field: str) -> str:
    tokens = split_name(field)
    translated = [TOKEN_TRANSLATIONS.get(token, token) for token in tokens]
    return " ".join(translated)


def description_for(entity: str, field: str, data_type: str) -> tuple[str, bool]:
    label = ENTITY_LABELS[entity]
    key = canonical_name(field)
    if key in EXACT_DESCRIPTIONS:
        return EXACT_DESCRIPTIONS[key].format(entity=label), False

    if key.endswith("_id"):
        stem = key[:-3]
        target = REFERENCE_LABELS.get(stem)
        if target:
            return f"Mã tham chiếu đến {target} liên quan đến {label}.", False
        return f"Mã tham chiếu đến {humanize_field(stem)} liên quan đến {label}.", True

    if key.endswith("_ids"):
        stem = key[:-4]
        target = REFERENCE_LABELS.get(stem, humanize_field(stem))
        return f"Danh sách mã {target} thuộc phạm vi của {label}.", stem not in REFERENCE_LABELS

    if key.startswith("is_"):
        phrase = humanize_field(key[3:])
        return f"Cho biết {label} có {phrase} hay không.", True
    if key.startswith("has_"):
        phrase = humanize_field(key[4:])
        return f"Cho biết {label} đã có {phrase} hay chưa.", True
    if key.startswith("show_"):
        phrase = humanize_field(key[5:])
        return f"Cho biết có hiển thị {phrase} trên đầu ra của {label} hay không.", True
    if key.startswith("allow_"):
        phrase = humanize_field(key[6:])
        return f"Cho biết cấu hình có cho phép {phrase} hay không.", True
    if key.startswith("require_") or key.startswith("requires_"):
        prefix = "require_" if key.startswith("require_") else "requires_"
        phrase = humanize_field(key[len(prefix):])
        return f"Cho biết nghiệp vụ có bắt buộc {phrase} trước khi hoàn tất hay không.", True
    if key.endswith("_at"):
        phrase = humanize_field(key[:-3])
        return f"Thời điểm {phrase} của {label}.", True
    if key.endswith("_date"):
        phrase = humanize_field(key[:-5])
        return f"Ngày {phrase} của {label}.", True
    if key.endswith("_time"):
        phrase = humanize_field(key[:-5])
        return f"Thời gian {phrase} của {label}.", True
    if key.endswith("_count"):
        phrase = humanize_field(key[:-6])
        return f"Số lượng {phrase} được ghi nhận cho {label}.", True
    if key.endswith("_quantity"):
        phrase = humanize_field(key[:-9])
        return f"Số lượng {phrase} của dòng nghiệp vụ.", True
    if key.endswith("_amount") or key.endswith("_money"):
        suffix = "_amount" if key.endswith("_amount") else "_money"
        phrase = humanize_field(key[: -len(suffix)])
        return f"Số tiền {phrase} được ghi nhận cho {label}.", True
    if key.endswith("_url"):
        phrase = humanize_field(key[:-4])
        return f"Đường dẫn truy cập {phrase} của {label}.", True
    if key.endswith("_urls"):
        phrase = humanize_field(key[:-5])
        return f"Danh sách đường dẫn truy cập {phrase} của {label}.", True
    if key.endswith("_name"):
        phrase = humanize_field(key[:-5])
        return f"Tên {phrase} được lưu để hiển thị và tra cứu {label}.", True
    if key.endswith("_code"):
        phrase = humanize_field(key[:-5])
        return f"Mã {phrase} dùng để nhận biết hoặc đối soát {label}.", True
    if key.endswith("_status"):
        phrase = humanize_field(key[:-7])
        return f"Trạng thái {phrase} tại thời điểm ghi nhận {label}.", True
    if key.endswith("_type") or key == "type":
        stem = key[:-5] if key.endswith("_type") else ""
        phrase = humanize_field(stem) if stem else label
        return f"Loại {phrase}, dùng để chọn quy tắc xử lý phù hợp.", True
    if key.endswith("_number"):
        phrase = humanize_field(key[:-7])
        return f"Số {phrase} hiển thị hoặc nhận từ hệ thống nguồn để tra cứu {label}.", True
    if key.endswith("_hash") or key.endswith("_fingerprint"):
        phrase = humanize_field(key.rsplit("_", 1)[0])
        return f"Mã băm của {phrase}, dùng để phát hiện nội dung thay đổi hoặc yêu cầu trùng.", True

    phrase = humanize_field(key)
    array_hint = "Danh sách" if "[]" in data_type or data_type.strip().startswith("Array<") else "Nội dung"
    return f"{array_hint} {phrase} được lưu cho {label} để phục vụ xử lý và tra cứu nghiệp vụ.", True


def iter_body_elements(doc: Document):
    from docx.table import Table

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = "Normal"
    run = new_para.add_run(text)
    run.bold = False
    fmt = new_para.paragraph_format
    fmt.space_before = paragraph.paragraph_format.space_before
    fmt.keep_with_next = True
    return new_para


def set_cell_text_preserve_style(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def mark_repeat_header(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def update_document(input_path: Path, output_path: Path, dry_run: bool) -> None:
    doc = Document(input_path)
    entity_headings = [
        p for p in doc.paragraphs if p.text.strip().startswith("Bảng (Entity):")
    ]
    entity_names = [p.text.split(":", 1)[1].strip() for p in entity_headings]

    if len(entity_names) != len(doc.tables):
        raise RuntimeError(
            f"Không thể ánh xạ tiêu đề và bảng: {len(entity_names)} tiêu đề, {len(doc.tables)} bảng."
        )
    missing_purposes = [name for name in entity_names if name not in ENTITY_PURPOSES]
    if missing_purposes:
        raise RuntimeError(f"Thiếu mô tả mục đích cho: {missing_purposes}")

    fallback_rows = []
    for entity, table in zip(entity_names, doc.tables):
        for row in table.rows[1:]:
            field = row.cells[0].text.strip()
            data_type = row.cells[1].text.strip()
            description, used_fallback = description_for(entity, field, data_type)
            if used_fallback:
                fallback_rows.append((entity, field, description))
            if not dry_run:
                set_cell_text_preserve_style(row.cells[2], description)
        if not dry_run:
            mark_repeat_header(table)

    if dry_run:
        print(f"Entities: {len(entity_names)}")
        print(f"Rows using generated fallback: {len(fallback_rows)}")
        for entity, field, description in fallback_rows:
            print(f"{entity}.{field}: {description}")
        return

    for paragraph in reversed(entity_headings):
        entity = paragraph.text.split(":", 1)[1].strip()
        insert_paragraph_after(paragraph, ENTITY_PURPOSES[entity])

    doc.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Updated entities: {len(entity_names)}")
    print(f"Updated attribute descriptions: {sum(len(t.rows) - 1 for t in doc.tables)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.dry_run and args.output is None:
        parser.error("output is required unless --dry-run is used")
    update_document(args.input, args.output or args.input, args.dry_run)


if __name__ == "__main__":
    main()
