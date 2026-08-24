from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Github\bduck-system")
OUT_DIR = ROOT / "outputs" / "invoice-management-guide"
SHOT_DIR = OUT_DIR / "screenshots"
OUTPUT = OUT_DIR / "HUONG_DAN_SU_DUNG_INVOICE_MANAGEMENT.docx"

NAVY = "0F2747"
BLUE = "246B9E"
LIGHT_BLUE = "E8F2FA"
INK = "172033"
MUTED = "667085"
LIGHT = "F5F7FA"
BORDER = "D9E0E8"
GREEN = "087A55"
LIGHT_GREEN = "E9F7F1"
AMBER = "9A6700"
LIGHT_AMBER = "FFF6D8"
RED = "B42318"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border(paragraph, color=BLUE, size="14", side="left"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "8")
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_note(doc, title, text, tone="info"):
    colors = {
        "info": (BLUE, LIGHT_BLUE),
        "warning": (AMBER, LIGHT_AMBER),
        "danger": (RED, LIGHT_RED),
        "success": (GREEN, LIGHT_GREEN),
    }
    accent, fill = colors[tone]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.12
    set_paragraph_border(p, accent)
    shade_paragraph(p, fill)
    r = p.add_run(f"{title}: ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=11, color=INK, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=11, color=INK, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK, italic=italic)
    return p


def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Bước {number}. {title}")
    set_run_font(r, size=11.5, color=NAVY, bold=True)
    p2 = doc.add_paragraph(style="Normal")
    p2.paragraph_format.left_indent = Inches(0.18)
    set_paragraph_border(p2, "B9C9D8", size="8")
    r2 = p2.add_run(text)
    set_run_font(r2, size=11, color=INK)
    return p2


def add_figure(doc, filename, caption, alt_text, width_inches=6.3, page_break_before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run()
    shape = run.add_picture(str(SHOT_DIR / filename), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return shape


def add_bullet(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.insert(0, num_pr)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def create_numbering(doc, abstract_id, num_id, fmt, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(value)
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row_idx, row in enumerate(rows):
        data_row = table.add_row()
        set_cant_split(data_row)
        cells = data_row.cells
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "F9FAFB")
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Trang ")
    set_run_font(r, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 30, NAVY, 0, 8),
    ("Subtitle", 14, MUTED, 0, 12),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, NAVY, 10, 5),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

styles["Caption"].font.name = "Calibri"
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)

bullet_id = create_numbering(doc, 41, 41, "bullet", "•", left=540, hanging=270)

header = section.header
hp = header.paragraphs[0]
hp.text = "J-PULSE  |  HƯỚNG DẪN NGHIỆP VỤ"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
set_run_font(hp.runs[0], size=8.5, color=MUTED, bold=True)
add_page_field(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.add_run("\n\n\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("SỔ TAY THAO TÁC")
set_run_font(r, size=10, color=BLUE, bold=True)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Quản lý hóa đơn điện tử")
set_run_font(r, size=30, color=NAVY, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Chức năng /invoice-management")
set_run_font(r, size=17, color=BLUE, bold=True)

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Đồng bộ • Kiểm tra trạng thái • Chỉnh sửa • Gửi lại an toàn")
set_run_font(r, size=13, color=MUTED)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(34)
p.paragraph_format.space_after = Pt(10)
set_paragraph_border(p, BLUE, size="18", side="top")

meta = add_table(
    doc,
    ["PHẠM VI", "ĐỐI TƯỢNG", "CẬP NHẬT"],
    [["J-PULSE / MISA meInvoice", "Kế toán, vận hành, quản trị", "23/08/2026"]],
    [3120, 3120, 3120],
    header_fill=NAVY,
    font_size=9.5,
)
for cell in meta.rows[0].cells:
    for run in cell.paragraphs[0].runs:
        set_run_font(run, size=9.5, color=WHITE, bold=True)

add_note(
    doc,
    "Nguyên tắc quan trọng nhất",
    "Không bấm phát hành hoặc gửi lại khi trạng thái đang chờ MISA, đang gửi MISA, đang đối soát hoặc chưa xác định kết quả. Luôn để hệ thống kiểm tra RefID để tránh tạo hóa đơn trùng.",
    "danger",
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(32)
r = p.add_run("TÀI LIỆU NỘI BỘ")
set_run_font(r, size=9, color=MUTED, bold=True)
p.add_run().add_break(WD_BREAK.PAGE)

# Front matter
doc.add_heading("Cách dùng tài liệu này", level=1)
add_body(doc, "Tài liệu mô tả giao diện đang vận hành của chức năng Quản lý hóa đơn, giải thích đầy đủ các trạng thái nghiệp vụ và hướng dẫn ba luồng thường dùng: kiểm tra hóa đơn, chỉnh sửa thông tin và gửi lại hóa đơn bị MISA từ chối.")
add_note(doc, "Dữ liệu minh họa", "Ảnh chụp sử dụng dữ liệu giao dịch mẫu trên hệ thống. Số lượng, mã đơn và quyền hiển thị có thể khác theo cửa hàng, ngày và tài khoản đăng nhập.", "info")

doc.add_heading("Mục lục", level=2)
for item in (
    "1. Điều kiện và quyền sử dụng",
    "2. Tổng quan giao diện và quy trình làm việc",
    "3. Giải thích trạng thái hóa đơn",
    "4. Chỉnh sửa thông tin hóa đơn",
    "5. Gửi lại hóa đơn bị từ chối",
    "6. Xử lý case sai lệch và hóa đơn đã phát hành",
    "7. Xử lý lỗi thường gặp",
    "8. Checklist vận hành hằng ngày",
):
    add_bullet(doc, item, bullet_id)

doc.add_heading("Quy ước cảnh báo", level=2)
add_table(
    doc,
    ["Mức", "Ý nghĩa", "Cách xử lý"],
    [
        ["Xanh / Thành công", "Dữ liệu đã hợp lệ hoặc hóa đơn đã phát hành.", "Tiếp tục theo đúng nút hành động đang bật."],
        ["Xanh dương / Đang xử lý", "Yêu cầu đã vào luồng hoặc đang chờ kết quả.", "Theo dõi; không tạo yêu cầu mới."],
        ["Vàng / Cần chú ý", "Thiếu dữ liệu, lỗi tạm thời hoặc cần kiểm tra.", "Đọc mô tả trạng thái trước khi thao tác."],
        ["Đỏ / Nguy cơ trùng", "Kết quả MISA chưa chắc chắn hoặc không cho phép phát hành lại.", "Dừng thao tác; đối soát RefID/TransactionID."],
    ],
    [1700, 3650, 4010],
)

# Section 1
doc.add_heading("1. Điều kiện và quyền sử dụng", level=1)
add_body(doc, "Đường dẫn: /invoice-management. Truy cập từ menu Quản lý hóa đơn trên thanh điều hướng bên trái.")
doc.add_heading("1.1. Điều kiện trước khi thao tác", level=2)
for text in (
    "Tài khoản được cấp phạm vi cửa hàng phù hợp.",
    "Cửa hàng đã cấu hình tài khoản MISA, ký hiệu hóa đơn, loại ký và thuế suất.",
    "Đơn hàng của ngày cần xử lý đã được cập nhật vào hệ thống.",
    "Người dùng có phương thức OTP hoạt động: ứng dụng xác thực hoặc email.",
):
    add_bullet(doc, text, bullet_id)

doc.add_heading("1.2. Quyền ảnh hưởng đến nút thao tác", level=2)
add_table(
    doc,
    ["Nhóm quyền", "Cho phép", "Biểu hiện khi thiếu quyền"],
    [
        ["Chuẩn bị hóa đơn", "Tạo/cập nhật draft, chỉnh thông tin.", "Nút tạo draft, cập nhật hoặc chỉnh sửa bị ẩn/vô hiệu."],
        ["Phát hành hàng loạt", "Chọn và phát hành hóa đơn đủ điều kiện.", "Checkbox hoặc nút Phát hành hóa đơn bị vô hiệu."],
        ["Thử lại", "Gửi lại đúng hóa đơn bị MISA từ chối.", "Checkbox case và nút Gửi lại bị vô hiệu."],
        ["Đối soát", "Đối chiếu MISA, làm mới case, đóng case.", "Không thể bấm Đối chiếu lại MISA hoặc Đóng case."],
        ["Cấu hình", "Mở và thay đổi cấu hình hóa đơn cửa hàng.", "Nút cấu hình không xuất hiện hoặc chỉ xem."],
    ],
    [1900, 3680, 3780],
)
add_note(doc, "Nếu nút bị mờ", "Kiểm tra cả quyền và trạng thái nghiệp vụ. Có quyền nhưng hóa đơn không đủ điều kiện thì nút vẫn bị vô hiệu để bảo vệ dữ liệu.", "warning")

# Section 2
doc.add_heading("2. Tổng quan giao diện và quy trình làm việc", level=1)
add_figure(
    doc,
    "01-overview-highlight.png",
    "Hình 1. Màn hình Quản lý hóa đơn và các vùng thao tác chính.",
    "Màn hình quản lý hóa đơn với các vùng được đánh số: bộ lọc cửa hàng và ngày, cập nhật dữ liệu, phát hành hóa đơn, các tab trạng thái và bộ lọc lỗi.",
    width_inches=5.95,
)
add_table(
    doc,
    ["Số", "Vùng thao tác", "Cách dùng"],
    [
        ["1", "Cửa hàng + khoảng ngày", "Chọn đúng cửa hàng và ngày/khoảng ngày trước mọi thao tác."],
        ["2", "Cập nhật dữ liệu", "Tải đơn nguồn, tính lại số tiền/thuế và đối chiếu trạng thái."],
        ["3", "Phát hành hóa đơn", "Mở luồng phát hành cho các hóa đơn đã chọn và đủ điều kiện."],
        ["4", "Nhóm trạng thái", "Cần xử lý, Sẵn sàng phát hành, Đã phát hành."],
        ["5", "Bộ lọc lỗi", "Lọc nhanh thiếu thông tin, sai tiền/thuế, chờ MISA, MISA từ chối, kiểm tra thủ công."],
    ],
    [700, 2600, 6060],
)

doc.add_heading("2.1. Quy trình chuẩn đầu ngày", level=2)
add_step(doc, 1, "Chọn phạm vi", "Chọn cửa hàng và ngày giao dịch. Với xử lý lịch sử, xác nhận kỹ khoảng ngày để tránh nhầm dữ liệu.")
add_step(doc, 2, "Cập nhật dữ liệu", "Bấm Cập nhật dữ liệu. Chờ thông báo hoàn tất và các thẻ Tổng đơn, Sẵn sàng phát hành, Cần xử lý, Đang gửi MISA được cập nhật.")
add_step(doc, 3, "Xử lý lỗi trước", "Mở Cần xử lý. Dùng các bộ lọc lỗi để sửa dữ liệu hoặc chuyển sang trang đối chiếu. Không bỏ qua trạng thái chờ MISA/kiểm tra thủ công.")
add_step(doc, 4, "Kiểm tra hóa đơn sẵn sàng", "Mở Sẵn sàng phát hành, vào Chi tiết kiểm tra, xác nhận người mua, hàng hóa, thuế suất, trước thuế, VAT và tổng tiền.")
add_step(doc, 5, "Phát hành và theo dõi", "Chỉ chọn hóa đơn đủ điều kiện, hoàn tất màn hình xác nhận và OTP. Theo dõi cho đến khi chuyển sang Đã phát hành trên MISA.")

# Section 3
doc.add_heading("3. Giải thích trạng thái hóa đơn", level=1)
add_body(doc, "Một hóa đơn đi qua ba giai đoạn: chuẩn bị dữ liệu, gửi sang MISA và đối soát sau gửi. Tên trên thẻ/tabs là tên thân thiện; phần Chi tiết kiểm tra hiển thị mô tả và hành động an toàn tương ứng.")

doc.add_heading("3.1. Trạng thái trước khi phát hành", level=2)
add_table(
    doc,
    ["Trạng thái", "Ý nghĩa", "Việc cần làm"],
    [
        ["Đã đồng bộ đơn nguồn", "Đơn đã tải về nhưng chưa có dữ liệu hóa đơn hoàn chỉnh.", "Tạo/chuẩn bị draft rồi kiểm tra."],
        ["Thiếu cấu hình thuế", "Thiếu thuế suất hoặc ánh xạ thuế của cửa hàng/sản phẩm.", "Bổ sung cấu hình thuế và cập nhật lại."],
        ["Cần chỉnh dữ liệu hóa đơn", "Người mua, dòng hàng, số lượng, đơn giá hoặc số tiền chưa hợp lệ.", "Mở Chi tiết kiểm tra và chỉnh sửa."],
        ["Sẵn sàng phát hành", "Dữ liệu đã vượt qua kiểm tra trước phát hành.", "Có thể chọn để gửi sang MISA sau khi rà soát."],
        ["Bản nháp đã bị từ chối", "Bị từ chối trong bước kiểm duyệt nội bộ; chưa gửi MISA.", "Chỉnh dữ liệu rồi chuẩn bị lại."],
    ],
    [2320, 3560, 3480],
)

doc.add_heading("3.2. Trạng thái trong khi gửi MISA", level=2)
add_table(
    doc,
    ["Trạng thái", "Ý nghĩa", "Việc cần làm"],
    [
        ["Đã xếp hàng — chưa gửi MISA", "Yêu cầu đang chờ đến lượt theo ký hiệu hóa đơn.", "Không phát hành lại."],
        ["Đang gửi dữ liệu sang MISA", "Backend đang xử lý yêu cầu; chưa có kết quả cuối.", "Không đóng luồng hoặc tạo yêu cầu mới."],
        ["MISA đã trả số — đang đối soát", "Đã có số hóa đơn/TransactionID nhưng trạng thái nội bộ chưa hoàn tất.", "Không gửi lại; chờ chuyển sang Đã phát hành."],
        ["Chưa nhận kết quả phát hành", "Chưa có số/TransactionID; hệ thống đang kiểm tra RefID.", "Tuyệt đối không gửi lại cho đến khi đối soát."],
        ["Lỗi kết nối — hệ thống sẽ thử lại", "Lỗi tạm thời khi gọi MISA.", "Không thử lại thủ công; theo dõi hệ thống tự xử lý."],
    ],
    [2320, 3560, 3480],
)

doc.add_page_break()
doc.add_heading("3.3. Trạng thái hoàn tất và ngoại lệ", level=2)
add_table(
    doc,
    ["Trạng thái", "Ý nghĩa", "Việc cần làm"],
    [
        ["Đã phát hành trên MISA", "MISA xác nhận phát hành; có thể đã có số hóa đơn và mã CQT.", "Xem/tải chứng từ; không phát hành lại."],
        ["MISA từ chối — có thể thử lại", "MISA từ chối trước khi tạo số; hệ thống xác định đủ điều kiện retry.", "Sửa nguyên nhân rồi dùng Gửi lại hóa đơn + OTP."],
        ["Cần đối soát thủ công — không được xuất lại", "Không thể xác định chắc MISA đã phát hành hay chưa.", "Kiểm tra RefID/TransactionID; không gửi lại."],
        ["Đã phát hành — cần kiểm tra sau phát hành", "Đã có hóa đơn nhưng còn sai lệch hoặc trạng thái CQT cần kiểm tra.", "Đối soát dữ liệu/CQT; không lập bản mới."],
        ["Đã hủy xử lý", "Yêu cầu phát hành đã bị hủy trước khi hoàn tất.", "Kiểm tra lý do và trạng thái MISA trước khi xử lý tiếp."],
        ["Đã đóng hồ sơ", "Hồ sơ đã xử lý xong và đóng.", "Không cần thao tác thêm."],
        ["Đã xóa/hủy trên MISA", "MISA ghi nhận hóa đơn bị xóa hoặc hủy.", "Kiểm tra lịch sử trước khi lập hóa đơn thay thế."],
        ["CQT từ chối / gửi CQT lỗi", "Hóa đơn có thể đã tồn tại nhưng cơ quan thuế chưa chấp nhận.", "Kiểm tra trạng thái CQT; không dùng retry phát hành gốc."],
    ],
    [2320, 3560, 3480],
)

add_figure(
    doc,
    "02-status-detail-highlight.png",
    "Hình 2. Cảnh báo đỏ: phải đối soát, không được xuất lại.",
    "Chi tiết hóa đơn có cảnh báo cần đối soát thủ công và nút cập nhật draft từ dữ liệu nguồn.",
    page_break_before=True,
)
add_body(doc, "Chú thích Hình 2: (1) cảnh báo kết quả MISA chưa chắc chắn; (2) chỉ cập nhật draft từ HKAPI khi dữ liệu nguồn thay đổi. Cập nhật draft không có nghĩa là được phép gửi lại.")

add_note(doc, "Quy tắc dừng", "Nếu có các cụm từ “không được xuất lại”, “đang gửi”, “đã xếp hàng”, “đang đối soát” hoặc “chưa nhận kết quả”, dừng mọi thao tác phát hành/gửi lại và chờ hoặc thực hiện đối soát.", "danger")

# Section 4
doc.add_heading("4. Chỉnh sửa thông tin hóa đơn", level=1)
add_body(doc, "Chỉnh sửa được dùng khi draft chưa phát hành và trạng thái cho phép. Sau phát hành, không sửa trực tiếp hóa đơn gốc; cần quy trình điều chỉnh/thay thế theo nghiệp vụ kế toán.")

doc.add_heading("4.1. Điều kiện được chỉnh sửa", level=2)
for text in (
    "Hóa đơn chưa phát hành và chưa ở trạng thái đang gửi/chờ xác nhận.",
    "Tài khoản có quyền chuẩn bị hóa đơn trong đúng cửa hàng.",
    "Nếu dữ liệu nguồn thay đổi, cập nhật draft từ HKAPI trước rồi kiểm tra revision mới.",
    "Nút Chỉnh sửa phải đang bật; nếu bị mờ, đọc cảnh báo trạng thái ở ngay phía trên.",
):
    add_bullet(doc, text, bullet_id)

doc.add_heading("4.2. Các bước thao tác", level=2)
add_step(doc, 1, "Mở đúng hóa đơn", "Chọn Cửa hàng, ngày và nhóm trạng thái. Tìm theo mã đơn/khách hàng, sau đó bấm vào dòng hóa đơn hoặc Chi tiết kiểm tra.")
add_step(doc, 2, "Kiểm tra cảnh báo", "Đọc Điểm cần xử lý và trạng thái trong phần Thông tin hóa đơn. Nếu draft cũ, bấm Cập nhật draft từ HKAPI rồi rà soát lại số tiền.")
add_step(doc, 3, "Vào chế độ chỉnh sửa", "Bấm Chỉnh sửa. Không thao tác nếu nút bị vô hiệu hoặc có cảnh báo không được xuất lại.")
add_step(doc, 4, "Cập nhật từng trường", "Sửa thông tin người mua, phương thức thanh toán và dòng hóa đơn. Giữ đúng chứng từ nguồn; không tự ý thay số tiền để làm cho hóa đơn vượt kiểm tra.")
add_step(doc, 5, "Hoàn tất", "Chờ trạng thái Đã tự động lưu, kiểm tra revision/số tiền đã tính lại, rồi bấm Xong. Mở lại Chi tiết kiểm tra để xác nhận trạng thái mới.")

add_figure(
    doc,
    "03-edit-form-highlight.png",
    "Hình 3. Biểu mẫu chỉnh sửa thông tin người mua và dòng hóa đơn.",
    "Biểu mẫu chỉnh sửa với vùng thông tin người mua, phương thức thanh toán, dòng hóa đơn và nút Xong được highlight.",
    page_break_before=True,
)
add_table(
    doc,
    ["Số", "Vùng chỉnh sửa", "Lưu ý"],
    [
        ["1", "Thông tin người mua", "Tên người mua/đơn vị, MST, địa chỉ, điện thoại, email phải khớp yêu cầu xuất hóa đơn."],
        ["2", "Phương thức thanh toán", "Dùng tên phương thức đúng cấu hình MISA, ví dụ TM/CK."],
        ["3", "Dòng hóa đơn", "Kiểm tra mã/tên hàng, ĐVT, SL, đơn giá, chiết khấu, VAT. Tổng được backend tính lại."],
        ["4", "Xong", "Chỉ bấm sau khi hiển thị Đã tự động lưu và đã kiểm tra lại số tiền."],
    ],
    [700, 2500, 6160],
)

doc.add_heading("4.3. Kiểm tra sau chỉnh sửa", level=2)
for text in (
    "Revision tăng hoặc thông tin lưu mới được hiển thị.",
    "Trước thuế + VAT = Tổng tiền; tổng tiền phải khớp chứng từ nguồn.",
    "Không còn lỗi validation trong Điểm cần xử lý.",
    "Trạng thái chuyển sang Sẵn sàng phát hành nếu mọi điều kiện hợp lệ.",
    "Nếu có chỉnh sửa tài chính, bảo đảm có căn cứ nghiệp vụ vì hệ thống ghi audit.",
):
    add_bullet(doc, text, bullet_id)

add_note(doc, "Không sửa sau phát hành", "Hóa đơn đã phát hành không được quay lại sửa draft để phát hành lần nữa. Liên hệ người phụ trách để dùng quy trình điều chỉnh/thay thế khi cần.", "danger")

# Section 5
doc.add_heading("5. Gửi lại hóa đơn bị MISA từ chối", level=1)
add_body(doc, "Gửi lại chỉ dành cho hóa đơn mà MISA đã từ chối rõ ràng trước khi tạo số hóa đơn và hệ thống đánh dấu Có thể gửi lại. Đây không phải nút khắc phục chung cho mọi lỗi.")

doc.add_heading("5.1. Khi nào được gửi lại", level=2)
add_table(
    doc,
    ["Tình huống", "Được gửi lại?", "Lý do"],
    [
        ["MISA từ chối, chưa có số/TransactionID, có nhãn Có thể gửi lại", "Có", "Hệ thống xác nhận retry eligible và giữ RefID chống trùng."],
        ["Đã xếp hàng / đang gửi MISA", "Không", "Yêu cầu cũ vẫn đang xử lý."],
        ["Chờ xác nhận hoặc chưa nhận kết quả", "Không", "Chưa biết MISA đã tạo hóa đơn hay chưa."],
        ["Cần đối soát thủ công", "Không", "Phải kiểm tra RefID/TransactionID trước."],
        ["Đã phát hành / có số hóa đơn", "Không", "Gửi lại sẽ có nguy cơ tạo trùng; dùng đối soát hoặc điều chỉnh."],
        ["Lỗi kết nối tạm thời", "Không thủ công", "Hệ thống tự thử lại."],
    ],
    [3500, 1500, 4360],
)

doc.add_heading("5.2. Các bước gửi lại an toàn", level=2)
add_step(doc, 1, "Mở trang xử lý sai lệch", "Tại tab Cần xử lý, bấm Xem trang đối chiếu. Trang Xử lý sai lệch hóa đơn sẽ mở theo cửa hàng và ngày đang chọn.")
add_step(doc, 2, "Lấy trạng thái mới nhất", "Chọn đúng cửa hàng/ngày và bấm Đối chiếu lại MISA. Thao tác này chỉ cập nhật dữ liệu đối soát, không tự gửi lại hóa đơn.")
add_step(doc, 3, "Lọc case", "Dùng ô tìm kiếm và bộ lọc. Chỉ chọn dòng có nhãn xanh Có thể gửi lại; checkbox của case không đủ điều kiện sẽ bị khóa.")
add_step(doc, 4, "Chọn cách gửi lại", "Bấm Gửi lại hóa đơn trên một dòng hoặc chọn nhiều case rồi bấm Gửi lại đã chọn (N). Kiểm tra đúng số lượng N.")
add_step(doc, 5, "Xác thực OTP", "Hộp Xác nhận gửi lại hóa đơn xuất hiện. Nhập OTP 6 số từ ứng dụng xác thực hoặc email, sau đó bấm Xác nhận.")
add_step(doc, 6, "Theo dõi kết quả", "Hệ thống kiểm tra trực tiếp RefID trên MISA trước khi xếp hàng. Theo dõi trạng thái đến khi Đã phát hành hoặc xuất hiện case cần xử lý mới.")

add_figure(
    doc,
    "04-retry-page-highlight.png",
    "Hình 4. Trang gửi lại hóa đơn với cơ chế chống trùng.",
    "Trang xử lý sai lệch với nút đối chiếu lại MISA, quy tắc chống gửi trùng, bộ lọc và nút gửi lại đã chọn được highlight.",
)
add_table(
    doc,
    ["Số", "Thao tác", "Ý nghĩa"],
    [
        ["1", "Đối chiếu lại MISA", "Lấy trạng thái mới nhất trước khi quyết định gửi lại."],
        ["2", "Quy tắc chống gửi trùng", "Đọc kỹ: nút chỉ xuất hiện khi MISA đã từ chối và hệ thống sẽ kiểm tra RefID."],
        ["3", "Tìm kiếm + bộ lọc", "Thu hẹp case đang mở và loại sai lệch cần xử lý."],
        ["4", "Gửi lại đã chọn (N)", "Chỉ bật khi đã chọn ít nhất một case đủ điều kiện."],
    ],
    [700, 2800, 5860],
)

doc.add_heading("5.3. Kết quả có thể gặp", level=2)
add_table(
    doc,
    ["Kết quả", "Diễn giải", "Hành động"],
    [
        ["Đã đưa vào hàng đợi gửi lại", "Kiểm tra RefID an toàn và yêu cầu đã được xếp hàng.", "Không bấm lại; theo dõi trạng thái."],
        ["MISA đã có dấu vết phát hành", "Hệ thống chặn gửi lại để tránh trùng.", "Đối soát theo RefID/TransactionID."],
        ["OTP sai/hết hạn", "Xác thực không thành công.", "Lấy OTP mới và thử lại một lần."],
        ["Không còn đủ điều kiện", "Trạng thái đã thay đổi sau lần tải dữ liệu.", "Làm mới/đối chiếu lại rồi đánh giá lại."],
        ["Tiếp tục bị MISA từ chối", "Nguyên nhân dữ liệu/cấu hình chưa được xử lý.", "Đọc mã lỗi, sửa nguyên nhân trước khi retry tiếp."],
    ],
    [2500, 3450, 3410],
)
add_note(doc, "OTP là bước cuối", "Sau khi nhập OTP hợp lệ và bấm Xác nhận, hệ thống có thể tạo yêu cầu gửi thật sang MISA. Luôn kiểm tra cửa hàng, ngày, mã đơn và số lượng hóa đơn trước bước này.", "warning")

# Section 6
doc.add_heading("6. Xử lý case sai lệch và hóa đơn đã phát hành", level=1)
doc.add_heading("6.1. Các loại case sai lệch", level=2)
add_table(
    doc,
    ["Loại case", "Ý nghĩa", "Xử lý gợi ý"],
    [
        ["Chưa tìm thấy hóa đơn trên MISA", "Có đơn trong hệ thống nhưng lần đối chiếu gần nhất không thấy hóa đơn tương ứng.", "Kiểm tra thời gian, RefID và trạng thái phát hành; không tự retry nếu chưa có nhãn cho phép."],
        ["Hóa đơn MISA không có đơn nguồn", "MISA có hóa đơn nhưng chưa ghép được với JPOS/HKAPI.", "Kiểm tra mã đơn, cửa hàng, ngày và liên kết nguồn."],
        ["Dữ liệu hóa đơn không khớp", "Thông tin hoặc số tiền giữa sổ nội bộ và MISA khác nhau.", "So sánh người mua, dòng hàng, thuế và tổng tiền."],
        ["Trạng thái phát hành không khớp", "Có dấu vết trên MISA nhưng trạng thái nội bộ chưa đúng.", "Đối soát TransactionID/số hóa đơn; không phát hành lại."],
        ["Hóa đơn đã bị xóa trên MISA", "MISA báo hóa đơn đã xóa/hủy.", "Kiểm tra lịch sử và quy trình thay thế."],
        ["Cơ quan thuế từ chối", "Hóa đơn đã lên MISA nhưng CQT không chấp nhận.", "Xử lý theo quy trình CQT, không retry phát hành gốc."],
        ["Cần kiểm tra thủ công", "Hệ thống không đủ dữ liệu để kết luận.", "Người phụ trách kiểm tra MISA và ghi kết quả."],
    ],
    [2600, 3400, 3360],
)

doc.add_heading("6.2. Đóng case không đồng nghĩa gửi lại", level=2)
add_body(doc, "Nút Đóng case chỉ ghi nhận kết quả xử lý nội bộ; không gọi MISA và không gửi lại hóa đơn. Chỉ đóng khi đã kiểm tra xong và nhập ghi chú đủ rõ để người khác hiểu bằng chứng/kết quả.")
add_note(doc, "Ghi chú tốt", "Nêu nơi đã kiểm tra, RefID/TransactionID hoặc số hóa đơn, kết quả trên MISA, thời điểm và quyết định xử lý. Không ghi chung chung như “đã kiểm tra”.", "info")

doc.add_heading("6.3. Hóa đơn đã phát hành", level=2)
for text in (
    "Mở tab Đã phát hành và tìm theo mã đơn, số hóa đơn, người mua hoặc MST.",
    "Kiểm tra số hóa đơn, ký hiệu, ngày hóa đơn, TransactionID và trạng thái CQT.",
    "Dùng Xem hóa đơn để mở bản hiển thị MISA; dùng tải PDF/XML để lưu chứng từ khi nút khả dụng.",
    "Nếu đã phát hành nhưng dữ liệu không khớp, mở case đối soát; không tạo lại hóa đơn gốc.",
):
    add_bullet(doc, text, bullet_id)

# Section 7
doc.add_heading("7. Xử lý lỗi thường gặp", level=1)
add_table(
    doc,
    ["Hiện tượng", "Nguyên nhân thường gặp", "Cách xử lý"],
    [
        ["Không thấy cửa hàng", "Tài khoản chưa được cấp phạm vi.", "Liên hệ quản trị phân quyền cửa hàng."],
        ["Cập nhật dữ liệu bị mờ", "Thiếu quyền, chưa chọn cửa hàng hoặc hệ thống đang đồng bộ.", "Kiểm tra phạm vi, chờ tiến trình kết thúc rồi làm mới."],
        ["Hóa đơn không vào Sẵn sàng", "Thiếu thuế, buyer, dòng hàng hoặc draft cũ.", "Mở Chi tiết kiểm tra và xử lý đúng từng lỗi."],
        ["Chỉnh sửa bị mờ", "Đang gửi/chờ MISA, đã phát hành, cần đối soát hoặc thiếu quyền.", "Không cố thao tác; xử lý trạng thái/đối soát trước."],
        ["Gửi lại bị mờ", "Không có case retry eligible hoặc chưa chọn case.", "Chỉ chọn dòng có nhãn Có thể gửi lại."],
        ["MISA đã có dấu vết phát hành", "RefID đã tồn tại trên MISA.", "Dừng retry, đối soát và cập nhật trạng thái nội bộ."],
        ["OTP không nhận được", "Email chậm, dùng TOTP hoặc địa chỉ email chưa đúng.", "Bấm gửi lại OTP; kiểm tra spam; liên hệ quản trị nếu tiếp diễn."],
        ["Tổng tiền khác nguồn", "Thuế, chiết khấu, số lượng/đơn giá hoặc nguồn đã thay đổi.", "Cập nhật draft từ HKAPI và so sánh từng dòng."],
        ["Đã phát hành nhưng CQT từ chối", "Lỗi nghiệp vụ thuế sau phát hành.", "Xử lý theo trạng thái CQT; không gửi lại hóa đơn gốc."],
    ],
    [2500, 3300, 3560],
    font_size=9.0,
)

doc.add_heading("7.1. Thông tin cần gửi khi yêu cầu hỗ trợ", level=2)
for text in (
    "Tên cửa hàng và ngày giao dịch.",
    "Mã đơn nguồn/Local/HK; không gửi ảnh chứa OTP.",
    "Trạng thái đang hiển thị và mã lỗi MISA (nếu có).",
    "RefID, TransactionID, số hóa đơn (nếu đã có).",
    "Ảnh chụp Chi tiết kiểm tra và thời điểm xảy ra lỗi.",
):
    add_bullet(doc, text, bullet_id)

# Section 8
doc.add_heading("8. Checklist vận hành hằng ngày", level=1)
add_table(
    doc,
    ["☐", "Kiểm tra", "Đạt khi"],
    [
        ["☐", "Đúng cửa hàng và ngày", "Phạm vi ở đầu trang đúng ca/ngày cần xử lý."],
        ["☐", "Đã cập nhật dữ liệu", "KPI và danh sách không còn trạng thái đang tải."],
        ["☐", "Cần xử lý = 0 hoặc có chủ sở hữu", "Mỗi lỗi có người phụ trách/ghi chú rõ."],
        ["☐", "Đang gửi MISA được theo dõi", "Không có yêu cầu treo quá thời gian dự kiến."],
        ["☐", "Hóa đơn sẵn sàng đã rà soát", "Buyer, dòng hàng, thuế và tổng tiền đúng."],
        ["☐", "Không phát hành/gửi lại trùng", "Không thao tác trên queued/submitting/pending/manual."],
        ["☐", "Đã phát hành được đối soát", "Có TransactionID/số hóa đơn và trạng thái CQT phù hợp."],
        ["☐", "Case đóng có ghi chú", "Ghi chú đủ bằng chứng và quyết định xử lý."],
    ],
    [650, 3380, 5330],
)

add_note(doc, "Kết thúc ca", "Nếu còn hóa đơn ở trạng thái chưa xác định hoặc cần đối soát thủ công, bàn giao bằng mã đơn + RefID/TransactionID + trạng thái hiện tại; tuyệt đối không bàn giao bằng chỉ dẫn “thử gửi lại”.", "danger")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_border(p, BLUE, size="12", side="top")
r = p.add_run("HẾT TÀI LIỆU")
set_run_font(r, size=9, color=MUTED, bold=True)

doc.core_properties.title = "Hướng dẫn sử dụng chức năng Quản lý hóa đơn"
doc.core_properties.subject = "J-PULSE /invoice-management"
doc.core_properties.author = "J-PULSE Operations"
doc.core_properties.keywords = "invoice-management, MISA meInvoice, hóa đơn, đối soát, gửi lại"

OUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
